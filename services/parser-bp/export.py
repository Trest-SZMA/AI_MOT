# -*- coding: utf-8 -*-
"""Выгрузка результатов парсера в Excel с цветовой заливкой."""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from parser_core import ScanReport

FILL_OK = PatternFill("solid", fgColor="C6EFCE")        # зелёный
FILL_BAD = PatternFill("solid", fgColor="FFC7CE")       # красный
FILL_UNKNOWN = PatternFill("solid", fgColor="FFEB9C")   # жёлтый
FILL_EMPTY = PatternFill("solid", fgColor="E4E4E4")     # серый — пустая заготовка
FONT_OK = Font(color="006100")
FONT_BAD = Font(color="9C0006")
FONT_UNKNOWN = Font(color="9C5700")
FONT_EMPTY = Font(color="4D4D4D")

HEADERS = [
    "Название БП и версия",
    "Кол-во месяцев по строкам",
    "Кол-во месяцев по ячейке",
    "Чистая прибыль",
    "Статус",
    "Расхождение",
    "Лукойл",
    "ДСП",
    "Файл",
    "Лист",
    "Лист с прибылью",
    "Ячейка блока",
    "Строки блока",
    "Полный путь",
    "Скрытые листы",
]

WIDTHS = [46, 16, 16, 18, 14, 58, 9, 9, 44, 28, 28, 13, 16, 60, 34]
COL_PROFIT = 4      # к этому столбцу применяем денежный формат


def build_workbook(report: ScanReport, folder: str = "") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Проверка месяцев"

    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F6228")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"

    for res in report.rows:
        ws.append([
            res.bp_name,
            res.months_by_rows,
            None if res.months_by_cell is None else res.months_by_cell,
            res.net_profit,
            res.status,
            res.mismatch,
            res.has_lukoil,
            res.has_dsp,
            res.file_name,
            res.sheet_name,
            res.profit_sheet,
            res.block_cell,
            res.rows_detail,
            res.file_path,
            res.hidden_sheets,
        ])
        row = ws.max_row
        if res.status == "ОК":
            fill, font = FILL_OK, FONT_OK
        elif res.status == "Расхождение":
            fill, font = FILL_BAD, FONT_BAD
        elif res.status == "Не заполнен":
            fill, font = FILL_EMPTY, FONT_EMPTY
        else:
            fill, font = FILL_UNKNOWN, FONT_UNKNOWN
        for col in range(1, len(HEADERS) + 1):
            cell = ws.cell(row, col)
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(vertical="center", wrap_text=(col in (1, 6)))
            if col == COL_PROFIT:
                cell.number_format = "# ##0"

    for idx, width in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    if ws.max_row > 1:
        ref = f"A1:{get_column_letter(len(HEADERS))}{ws.max_row}"
        table = Table(displayName="Проверка", ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
        ws.add_table(table)

    _add_summary(wb, report, folder)
    _add_skipped(wb, report)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _add_summary(wb, report: ScanReport, folder: str):
    ws = wb.create_sheet("Сводка")
    ok = sum(1 for r in report.rows if r.status == "ОК")
    bad = sum(1 for r in report.rows if r.status == "Расхождение")
    empty = sum(1 for r in report.rows if r.status == "Не заполнен")
    unknown = len(report.rows) - ok - bad - empty
    data = [
        ("Папка", folder),
        ("Файлов обработано", report.files_scanned),
        ("Листов просмотрено", report.sheets_scanned),
        ("Найдено блоков «Расчет процентов»", len(report.rows)),
        ("Совпадает (зелёный)", ok),
        ("Расхождение (красный)", bad),
        ("Не заполнен (серый)", empty),
        ("Не определено (жёлтый)", unknown),
        ("Файлов пропущено", len(report.skipped)),
    ]
    for key, value in data:
        ws.append([key, value])
    for cell in ws["A"]:
        cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 80


def _add_skipped(wb, report: ScanReport):
    if not report.skipped:
        return
    ws = wb.create_sheet("Пропущенные файлы")
    ws.append(["Файл", "Причина"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for path, reason in report.skipped:
        ws.append([path, reason])
    ws.column_dimensions["A"].width = 80
    ws.column_dimensions["B"].width = 60
