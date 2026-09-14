#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Проверка выгрузки в 1С: тот же код, что в браузере, читается openpyxl.

1С не читает inlineStr, поэтому файл строится через sharedStrings (t="s") +
styles.xml + cellStyles, лист «Лист_1». Функция buildXlsx живёт в web/tabs.js;
здесь она запускается в Node с минимальными заглушками (без браузера), файл
пишется на диск и читается openpyxl — предупреждений быть не должно.

    python3 tools/check_1c_xlsx.py        (нужен node)
"""
import os, sys, json, subprocess, warnings, zipfile, tempfile

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TABS = os.path.join(HERE, "web", "tabs.js")
OUT = os.path.join(HERE, "out", "_test")
os.makedirs(OUT, exist_ok=True)

# Данные берём из реального расчёта, если он есть, иначе минимальный образец.
sd = os.path.join(HERE, "out", "sales_data.json")
rows_src = []
if os.path.exists(sd):
    D = json.load(open(sd, encoding="utf-8"))
    for it in D["items"][:200]:
        real = it["series"] != "БЕЗ СЕРИИ" and not str(it["series"]).startswith("площадка:")
        rows_src.append([it["sklad"], it["unit"], it["code"], it["base"], it["name"],
                         it["series_full"] if real else "", it["tonnes"]])
if not rows_src:
    rows_src = [["Склад МАТ Б.Осенцы", "т", "00000005139", "База Осенцы",
                 "Труба НКТ 73*5,5 б/у*", "1578 22С0371 от 11.03.22 г. (ЗС-УВМ) ДС35 V", 0.756]]

HARNESS = r"""
const fs = require('fs');
const src = fs.readFileSync(process.argv[2], 'utf8');
// Заглушки: tabs.js на верхнем уровне только объявляет функции и читает DATA.
global.DATA = {meta:{}, items:[], buys:[], flows:[], flows_var:[], moves:[], S:[], nm:{}};
global.META = {}; global.ITEMS = []; global.S = [];
global.str = x => (typeof x === 'number' ? '' : (x || ''));
global.esc = s => String(s == null ? '' : s);
global.fmt = n => String(n); global.fmt3 = n => String(n);
global.el = (t, c, h) => ({t, c, h, kids: [], appendChild(x){this.kids.push(x);},
                           querySelector(){return null;}, set innerHTML(v){this.h = v;}});
global.$ = () => null;
global.groupBy = (a, f) => { const m = new Map(); for (const x of a) {
  const k = f(x); if (!m.has(k)) m.set(k, []); m.get(k).push(x); } return m; };
global.sum = (a, f) => a.reduce((s, x) => s + (f(x) || 0), 0);
global.worstOf = () => null; global.worstBadge = a => a; global.badge = () => '';
global.BADGE_RANK = {}; global.BADGE_TITLE = {}; global.REAL = () => true;
global.baseSort = (a, b) => 0; global.TABS = {};
global.window = global; global.document = {querySelector: () => null,
  querySelectorAll: () => [], createElement: () => ({style:{}, addEventListener(){}}),
  documentElement: {dataset:{}}, body:{appendChild(){}}};
global.localStorage = {};
eval(src);
const payload = JSON.parse(fs.readFileSync(process.argv[3], 'utf8'));
for (const job of payload) {
  const bytes = buildXlsx('Лист_1', job.header, job.rows);
  fs.writeFileSync(job.out, Buffer.from(bytes));
}
console.log('node: собрано файлов ' + payload.length);
"""

def main():
    hpath = os.path.join(tempfile.gettempdir(), "xlsx_harness.js")
    open(hpath, "w", encoding="utf-8").write(HARNESS)
    jobs = [
        {"out": os.path.join(OUT, "1c_tpl.xlsx"),
         "header": ["Номенклатура", "Количество", "Количество(в единицах хранения)",
                    "Статус указания серий", "Статус указания серий отправитель",
                    "Статус указания серий получатель", "Серия"],
         "rows": [[r[4], r[6], r[6], 14, "", 14, r[5]] for r in rows_src]},
        {"out": os.path.join(OUT, "1c_ext.xlsx"),
         "header": ["Склад", "Ед. изм.", "Номенклатура.Код", "Подразделение",
                    "Номенклатура", "Серия", "Количество"],
         "rows": rows_src},
    ]
    jpath = os.path.join(tempfile.gettempdir(), "xlsx_jobs.json")
    json.dump(jobs, open(jpath, "w", encoding="utf-8"), ensure_ascii=False)
    try:
        out = subprocess.run(["node", hpath, TABS, jpath], capture_output=True, text=True)
    except FileNotFoundError:
        sys.exit("node не найден — проверку выгрузки в 1С запустить нечем")
    if out.returncode:
        sys.exit("node упал:\n" + out.stderr[:2000])
    print(out.stdout.strip())

    import openpyxl
    bad = 0
    for job in jobs:
        fn = job["out"]
        with warnings.catch_warnings(record=True) as w:
            warnings.simplefilter("always")
            wb = openpyxl.load_workbook(fn)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
        z = zipfile.ZipFile(fn)
        sheet = z.read("xl/worksheets/sheet1.xml").decode()
        styles = z.read("xl/styles.xml").decode()
        checks = [
            ("лист называется Лист_1", wb.sheetnames == ["Лист_1"]),
            ("нет inlineStr", "inlineStr" not in sheet),
            ("строки через sharedStrings (t=\"s\")", 't="s"' in sheet),
            ("есть styles.xml с cellStyles", "cellStyles" in styles),
            ("openpyxl читает без предупреждений", not w),
            ("шапка на месте", list(rows[0]) == job["header"]),
            ("числа остались числами",
             any(isinstance(v, (int, float)) for v in rows[1])),
            ("строк столько же, сколько позиций", ws.max_row == len(job["rows"]) + 1),
        ]
        print("\n%s (%.1f КБ)" % (os.path.basename(fn), os.path.getsize(fn) / 1024))
        for name, ok in checks:
            print("   [%s] %s" % ("✓" if ok else "✗", name))
            if not ok: bad += 1
        if w:
            print("   предупреждения:", [str(x.message) for x in w])
        print("   пример строки:", rows[1])
    print("\nИТОГ: " + ("выгрузка в 1С корректна" if not bad else "ПРОБЛЕМ: %d" % bad))
    sys.exit(0 if not bad else 1)


if __name__ == "__main__":
    main()
