"""Smoke-тесты API: экспорт и ключевые эндпоинты на чистом каталоге данных.

Регрессия, которую ловим: экспорт Excel падал с 500 на сервере, где каталог
data/exports ещё не существовал (создавался только при загрузке файла).
"""
import os
import zipfile

import pytest


@pytest.fixture()
def client(tmp_path, monkeypatch):
    """Приложение на пустом каталоге данных — как на свежем сервере."""
    monkeypatch.setenv("METOPTTORG_DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("METOPTTORG_AUTH", "off")  # тестовый режим без входа
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path / 'test.db'}")
    for mod in ("app.db.session", "app.api.main"):
        import sys
        sys.modules.pop(mod, None)
    from fastapi.testclient import TestClient
    from app.api.main import app
    from app.db.session import SessionLocal
    from app.db.models import KpDocument, KpPosition

    with TestClient(app) as c:  # startup создаёт каталоги и нормативы
        s = SessionLocal()
        doc = KpDocument(title="Тестовое КП")
        s.add(doc)
        s.commit()
        s.add(KpPosition(document_id=doc.id, raw_name="Труба НКТ 73 б/у",
                         quantity=10, unit="т", weight_kg=10_000))
        s.commit()
        s.close()
        yield c


def test_ui_served(client):
    r = client.get("/")
    assert r.status_code == 200
    assert "МетОптТорг" in r.text


def test_kp_calc_and_totals(client):
    r = client.get("/api/kp/1")
    assert r.status_code == 200
    d = r.json()
    assert len(d["positions"]) == 1
    assert d["totals"]["base_buyout"] >= 0
    assert d["priced_at"]


def test_export_on_fresh_data_dir(client):
    """Экспорт работает, даже если каталог экспорта ещё не создавался."""
    r = client.get("/api/kp/1/export")
    assert r.status_code == 200, r.text
    assert len(r.content) > 4000
    tmp = "/tmp/_export_smoke.xlsx"
    with open(tmp, "wb") as f:
        f.write(r.content)
    with zipfile.ZipFile(tmp) as z:  # открывается как настоящий xlsx
        assert any("sheet" in n for n in z.namelist())
    os.unlink(tmp)


def test_scrap_prices_and_norms(client):
    prices = client.get("/api/scrap-prices").json()
    assert {p["material"] for p in prices} >= {"медь", "чермет", "свинец", "16АЦ"}
    norms = {n["key"] for n in client.get("/api/norms").json()}
    assert "target_margin_percent" in norms
    assert client.get("/api/norms/regional").status_code == 200


def test_manual_patch_recalculates(client):
    before = client.get("/api/kp/1").json()["positions"][0]
    # без цены реализации итог не изменится: годное оценивается по лому (floor),
    # поэтому проверяем пересчёт с ценой — 50 000 ₽/т против лома 25 000 ₽/т
    r = client.patch("/api/kp/position/1",
                     json={"good_percent": 40, "resale_price": 50_000})
    assert r.status_code == 200
    after = r.json()
    assert after["base"]["good_percent"] == 40
    assert after["base"]["good_source"] == "ручная корректировка"
    assert after["base"]["good_price_source"] == "ручная корректировка"
    # лом 6 т × 25 000 + годное 4 т × 50 000 = 350 000 ₽
    assert after["base"]["resale_total"] == pytest.approx(350_000)
    assert after["base"]["resale_total"] > before["base"]["resale_total"]


def test_ai_status_reports_disabled_without_key(client, monkeypatch):
    monkeypatch.delenv("OPENROUTER_API_KEY", raising=False)
    monkeypatch.delenv("PERPLEXITY_API_KEY", raising=False)
    assert client.get("/api/ai-status").json() == {"enabled": False}


# ------------------------------------------------------------ авторизация


@pytest.fixture()
def auth_client(tmp_path, monkeypatch):
    """Приложение с ВКЛЮЧЁННОЙ авторизацией и двумя пользователями."""
    monkeypatch.setenv("METOPTTORG_DATA_DIR", str(tmp_path / "authdata"))
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path / 'auth.db'}")
    monkeypatch.delenv("METOPTTORG_AUTH", raising=False)
    import sys
    for mod in ("app.db.session", "app.api.main", "app.auth"):
        sys.modules.pop(mod, None)
    from fastapi.testclient import TestClient
    from app.api.main import app
    from app.auth import create_user
    from app.db.session import SessionLocal

    with TestClient(app) as c:
        s = SessionLocal()
        create_user(s, "director", "director-pass-123", "директор", "Директор")
        create_user(s, "buyer", "buyer-pass-123", "оценщик", "Оценщик")
        s.close()
        yield c


def test_anonymous_is_rejected(auth_client):
    r = auth_client.get("/api/kp")
    assert r.status_code == 401
    assert "Basic" in r.headers.get("WWW-Authenticate", "")


def test_wrong_password_rejected(auth_client):
    assert auth_client.get("/api/kp", auth=("director", "неверный")).status_code == 401
    assert auth_client.get("/api/kp", auth=("нет-такого", "x")).status_code == 401


def test_appraiser_can_work_but_not_approve(auth_client):
    buyer = ("buyer", "buyer-pass-123")
    assert auth_client.get("/api/kp", auth=buyer).status_code == 200
    assert auth_client.get("/api/me", auth=buyer).json()["can_approve"] is False
    # утверждение норматива — только директор
    r = auth_client.post("/api/norms", json={"key": "target_margin_percent",
                                             "value": 30}, auth=buyer)
    assert r.status_code == 403
    assert "директор" in r.json()["detail"]


def test_director_can_approve_and_is_recorded(auth_client):
    director = ("director", "director-pass-123")
    r = auth_client.post("/api/norms",
                         json={"key": "target_margin_percent", "value": 30},
                         auth=director)
    assert r.status_code == 200
    norms = {n["key"]: n for n in auth_client.get("/api/norms", auth=director).json()}
    assert norms["target_margin_percent"]["value"] == 30
    assert norms["target_margin_percent"]["approved_by"] == "director"


# --------------------------------------------- сверка тоннажа и документы


def test_tonnage_reconciliation(tmp_path, monkeypatch):
    """Итог из строки ИТОГО файла сверяется с разобранным тоннажем."""
    import openpyxl
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker
    from app.db.models import Base
    from app.importers.kp import import_kp

    f = tmp_path / "nvo.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Приложение № 1 к письму"])
    ws.append(["№", "Наименование ТМЦ", "Ед. изм.", "Кол-во", "Вес, т"])
    ws.append([1, "Труба НКТ 73 б/у", "т", 10, 10.0])
    ws.append([2, "Труба НКТ 89 б/у", "т", 5, 5.0])
    ws.append(["ИТОГО", None, None, None, 15.0])
    wb.save(f)

    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    s = sessionmaker(bind=engine)()
    doc, res = import_kp(s, str(f), "nvo.xlsx")
    assert res.file_total_t == pytest.approx(15.0)
    assert res.total_weight_t == pytest.approx(15.0)
    assert not res.warnings  # сошлось — предупреждений нет
    assert doc.file_total_weight_kg == pytest.approx(15_000)
    s.close()


def test_reconciliation_warns_on_mismatch(tmp_path):
    """Если часть строк не разобрана — импорт честно предупреждает."""
    import openpyxl
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker
    from app.db.models import Base
    from app.importers.kp import import_kp

    f = tmp_path / "nvo2.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["№", "Наименование ТМЦ", "Ед. изм.", "Кол-во", "Вес, т"])
    ws.append([1, "Труба НКТ 73 б/у", "т", 10, 10.0])
    ws.append(["ИТОГО", None, None, None, 42.0])  # в файле больше, чем строк
    wb.save(f)
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    s = sessionmaker(bind=engine)()
    _, res = import_kp(s, str(f), "nvo2.xlsx")
    assert res.warnings and "расхождение" in res.warnings[0]
    s.close()


def test_offer_form_renders(client):
    r = client.get("/api/kp/1/offer?scenario=base")
    assert r.status_code == 200
    assert "Ценовое предложение" in r.text
    assert "МетОптТорг" in r.text
    assert "Труба НКТ 73 б/у" in r.text
    assert client.get("/api/kp/1/offer?scenario=bp").status_code == 200
    assert client.get("/api/kp/1/offer?scenario=xxx").status_code == 422


def test_export_has_formatting_and_conditions_sheet(client, tmp_path):
    import openpyxl
    r = client.get("/api/kp/1/export")
    assert r.status_code == 200
    f = tmp_path / "e.xlsx"
    f.write_bytes(r.content)
    wb = openpyxl.load_workbook(f)
    assert wb.sheetnames == ["Оценка КП", "Условия расчёта"]
    ws = wb["Оценка КП"]
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref is not None
    assert ws["A1"].font.bold and ws.column_dimensions["A"].width > 20
    assert ws.cell(row=ws.max_row, column=1).value == "ИТОГО"
    keys = {row[0].value for row in wb["Условия расчёта"].iter_rows()}
    assert "target_margin_percent" in keys  # условия воспроизводимы


def test_schema_autoupgrade_adds_missing_columns(tmp_path, monkeypatch):
    """Новая колонка в модели добавляется в уже существующую базу.

    Регрессия: после добавления kp_documents.file_total_weight_kg сервер падал
    с «no such column», потому что create_all не меняет существующие таблицы.
    """
    import sqlite3
    import sys
    db = tmp_path / "old.db"
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{db}")
    monkeypatch.setenv("METOPTTORG_DATA_DIR", str(tmp_path / "d"))
    for mod in ("app.db.session",):
        sys.modules.pop(mod, None)
    from app.db.session import init_db
    init_db()
    # имитируем старую базу: удаляем колонку через пересоздание таблицы
    con = sqlite3.connect(db)
    con.execute("ALTER TABLE kp_documents DROP COLUMN file_total_weight_kg")
    con.commit()
    cols = {r[1] for r in con.execute("PRAGMA table_info(kp_documents)")}
    assert "file_total_weight_kg" not in cols
    con.close()

    from app.db.session import engine, ensure_schema
    engine.dispose()  # пул держит соединение со старой схемой
    added = ensure_schema()
    assert any("file_total_weight_kg" in a for a in added)
    con = sqlite3.connect(db)
    cols = {r[1] for r in con.execute("PRAGMA table_info(kp_documents)")}
    con.close()
    assert "file_total_weight_kg" in cols


# ------------------------------------------------- справочник номенклатуры


@pytest.fixture()
def catalog_client(client):
    """Справочник с дублем, выбросом и обычными позициями."""
    from app.calc.analog_price import reset_cache
    from app.db.models import Item, ItemAlias, MergeCandidate, PriceQuote
    from app.db.session import SessionLocal
    from app.normalize import normalize_name

    s = SessionLocal()
    def add(name, price=None, family="ПЭД", unit="шт", trade=True):
        it = Item(name=name, name_normalized=normalize_name(name), unit=unit,
                  family=family, is_trade=trade)
        s.add(it)
        s.commit()
        if price:
            s.add(PriceQuote(item_id=it.id, quote_type="sales_fact",
                             price=price, unit=unit, ttl_days=9999))
            s.commit()
        return it
    for i, price in enumerate((40_000, 45_000, 50_000, 55_000, 60_000)):
        add(f"Двигатель ПЭД-{i} б/у", price)
    outlier = add("Двигатель ПЭД-редкий б/у", 900_000)      # ×18 к медиане
    merged = add("Труба 89 б/у", 30_000, family="труба", unit="т")
    s.add(ItemAlias(item_id=merged.id, guid="dup-guid",
                    name="Труба 89 б/у (тн)",
                    name_normalized="труба 89 б/у"))
    a = add("Секция ЭЦН 5-50 б/у", family="секция ЭЦН")
    b = add("Секция ЭЦН 5 50 б/у", family="секция ЭЦН")
    s.add(MergeCandidate(item_id_a=a.id, item_id_b=b.id, score=1.0))
    s.commit()
    s.close()
    reset_cache()
    return client


def test_items_list_and_filters(catalog_client):
    d = catalog_client.get("/api/items").json()
    assert d["total"] >= 9
    names = {i["name"] for i in d["items"]}
    assert "Труба 89 б/у" in names

    merged = catalog_client.get("/api/items?flag=merged").json()["items"]
    assert [i["name"] for i in merged] == ["Труба 89 б/у"]
    assert merged[0]["merged_duplicates"] == 1

    no_price = catalog_client.get("/api/items?flag=no_price").json()["items"]
    assert all(i["sale_price"] is None for i in no_price)

    fam = catalog_client.get("/api/items?family=ПЭД").json()["items"]
    assert all(i["family"] == "ПЭД" for i in fam)


def test_price_outlier_detected(catalog_client):
    out = catalog_client.get("/api/items?flag=outlier").json()["items"]
    assert [i["name"] for i in out] == ["Двигатель ПЭД-редкий б/у"]
    assert out[0]["outlier_ratio"] > 3  # цена втрое и более выше медианы
    # обычные позиции выбросами не помечены
    normal = catalog_client.get("/api/items?q=ПЭД-1").json()["items"]
    assert all(i["outlier_ratio"] is None for i in normal)


def test_item_card_shows_prices_and_aliases(catalog_client):
    items = catalog_client.get("/api/items?flag=merged").json()["items"]
    card = catalog_client.get(f"/api/items/{items[0]['id']}").json()
    assert card["name"] == "Труба 89 б/у"
    assert [a["name"] for a in card["aliases"]] == ["Труба 89 б/у (тн)"]
    assert any(p["type"] == "sales_fact" for p in card["prices"])


def test_merge_candidate_merge_keeps_guid(catalog_client):
    cands = catalog_client.get("/api/merge-candidates").json()
    assert len(cands) == 1
    cand = cands[0]
    r = catalog_client.post(f"/api/merge-candidates/{cand['id']}/merge")
    assert r.status_code == 200
    assert catalog_client.get("/api/merge-candidates").json() == []
    card = catalog_client.get(f"/api/items/{cand['a']['id']}").json()
    assert cand["b"]["name"] in [a["name"] for a in card["aliases"]]
    # поглощённая карточка исчезла из списка торговых
    names = {i["name"] for i in catalog_client.get("/api/items").json()["items"]}
    assert cand["b"]["name"] not in names


def test_merge_candidate_reject(catalog_client):
    cand = catalog_client.get("/api/merge-candidates").json()[0]
    assert catalog_client.post(
        f"/api/merge-candidates/{cand['id']}/reject").status_code == 200
    assert catalog_client.get("/api/merge-candidates").json() == []
    # обе позиции остались
    names = {i["name"] for i in catalog_client.get("/api/items").json()["items"]}
    assert cand["a"]["name"] in names and cand["b"]["name"] in names


def test_item_card_shows_only_relevant_yields(catalog_client):
    """Регрессия: сравнение size_key с None давало IS NULL и притягивало
    нормативы чужих категорий в карточку позиции."""
    from app.db.models import ExpertYield
    from app.db.session import SessionLocal

    s = SessionLocal()
    s.add(ExpertYield(block="expert", category="труба", good_percent=50))
    s.add(ExpertYield(block="ai", category="ПЭД", good_percent=25))
    s.commit()
    s.close()
    items = catalog_client.get("/api/items?family=секция ЭЦН").json()["items"]
    card = catalog_client.get(f"/api/items/{items[0]['id']}").json()
    scopes = {y["scope"] for y in card["yields"]}
    assert "категория труба" not in scopes
    assert "категория ПЭД" not in scopes


def test_ai_price_refused_for_scrap_items(catalog_client, monkeypatch):
    """Для лома ИИ-поиск запрещён: он находит цену б/у изделия, а не лома."""
    from app.db.models import Item
    from app.db.session import SessionLocal
    from app.normalize import normalize_name

    monkeypatch.setenv("PERPLEXITY_API_KEY", "test-key")  # провайдер «есть»
    s = SessionLocal()
    it = Item(name="Лом 5А (Труба 159)", name_normalized=normalize_name("Лом 5А"),
              unit="т", family="лом", is_trade=True)
    s.add(it)
    s.commit()
    item_id = it.id
    s.close()
    r = catalog_client.post(f"/api/items/{item_id}/ai-price")
    assert r.status_code == 422
    assert "прайс" in r.json()["detail"]


# ---------------------------------------------------- пакетный ИИ-поиск цен


def test_batch_skips_scrap_and_already_priced(catalog_client, monkeypatch):
    """В пачку не попадают лом и позиции, у которых цена уже есть."""
    from app import ai_price
    from app.db.models import Item
    from app.db.session import SessionLocal
    from app.normalize import normalize_name

    s = SessionLocal()
    s.add(Item(name="Лом 12А (т)", name_normalized=normalize_name("Лом 12А"),
               unit="т", family="лом", is_trade=True))
    s.add(Item(name="Гидрозащита ПР92Д б/у", unit="шт", family="гидрозащита",
               name_normalized=normalize_name("Гидрозащита ПР92Д б/у"),
               is_trade=True))
    s.commit()
    s.close()

    started = {}

    def fake_start(factory, ids, user=""):
        started["ids"] = list(ids)
        return {"id": "test", "total": len(ids), "done": 0, "found": 0,
                "no_data": 0, "skipped": 0, "errors": 0, "status": "готово"}

    monkeypatch.setattr(ai_price, "start_batch", fake_start)
    monkeypatch.setenv("PERPLEXITY_API_KEY", "test-key")

    r = catalog_client.post("/api/ai-price/batch", json={"family": "лом"})
    assert r.status_code == 200
    assert r.json()["job"] is None  # ломовых кандидатов нет

    # все ПЭД из фикстуры уже с ценой продаж — кандидатов тоже нет
    assert catalog_client.post(
        "/api/ai-price/batch", json={"family": "ПЭД"}).json()["job"] is None

    # а гидрозащита без цены — попадает в пачку
    r = catalog_client.post("/api/ai-price/batch", json={"family": "гидрозащита"})
    assert r.status_code == 200 and r.json()["job"]["total"] == 1
    from app.db.session import SessionLocal as SL
    s = SL()
    names = {i.name for i in s.query(Item).filter(
        Item.id.in_(started["ids"])).all()}
    s.close()
    assert names == {"Гидрозащита ПР92Д б/у"}


def test_batch_requires_scope(catalog_client, monkeypatch):
    monkeypatch.setenv("PERPLEXITY_API_KEY", "test-key")
    r = catalog_client.post("/api/ai-price/batch", json={})
    assert r.status_code == 422


def test_batch_reports_progress(catalog_client, monkeypatch):
    """Прогресс доступен по идентификатору задачи; неизвестный — 404."""
    from app import ai_price

    monkeypatch.setattr(ai_price, "_jobs",
                        {"abc": {"id": "abc", "total": 3, "done": 2,
                                 "found": 1, "status": "работает"}})
    r = catalog_client.get("/api/ai-price/batch/abc")
    assert r.status_code == 200 and r.json()["done"] == 2
    assert catalog_client.get("/api/ai-price/batch/нет").status_code == 404


# ---------------------------------------------------- ИИ-оценка выхода


def test_ai_yield_refused_for_metal_only_families(catalog_client, monkeypatch):
    """Для лома и кабеля деловой выход не применяется — ценность в металле."""
    from app.db.models import Item
    from app.db.session import SessionLocal
    from app.normalize import normalize_name

    monkeypatch.setenv("PERPLEXITY_API_KEY", "test-key")
    s = SessionLocal()
    ids = {}
    for name, family in (("Лом 5А (т)", "лом"), ("Кабель ВВГ 3х4", "кабель")):
        it = Item(name=name, name_normalized=normalize_name(name), unit="т",
                  family=family, is_trade=True)
        s.add(it)
        s.commit()
        ids[family] = it.id
    s.close()
    for family, item_id in ids.items():
        r = catalog_client.post(f"/api/items/{item_id}/ai-yield")
        assert r.status_code == 422, family
        assert "металлосодержании" in r.json()["detail"]


def test_ai_yield_saved_as_ai_block_and_loses_to_expert(catalog_client, monkeypatch):
    """ИИ-выход пишется в блок ai: работает в БП и уступает эксперту."""
    from app import ai_yield
    from app.db.models import ExpertYield, Item
    from app.db.session import SessionLocal
    from app.normalize import normalize_name

    monkeypatch.setenv("PERPLEXITY_API_KEY", "test-key")
    monkeypatch.setattr(ai_yield, "ask_yield", lambda name, family: {
        "good_percent": 42, "confidence": "low",
        "sources": ["https://example.com/a"], "comment": "по объявлениям"})

    s = SessionLocal()
    it = Item(name="Гидрозащита ПР92Д б/у", unit="шт", family="гидрозащита",
              name_normalized=normalize_name("Гидрозащита ПР92Д б/у"),
              is_trade=True)
    s.add(it)
    s.commit()
    item_id = it.id
    s.close()

    r = catalog_client.post(f"/api/items/{item_id}/ai-yield")
    assert r.status_code == 200 and r.json()["good_percent"] == 42

    s = SessionLocal()
    row = s.query(ExpertYield).filter(ExpertYield.item_id == item_id).first()
    assert row.block == "ai" and row.confidence == "low"
    assert row.source_url == "https://example.com/a"
    s.close()

    # повторный запрос заменяет прежнюю ИИ-оценку, а не плодит дубли
    monkeypatch.setattr(ai_yield, "ask_yield", lambda name, family: {
        "good_percent": 55, "sources": [], "comment": "уточнено"})
    catalog_client.post(f"/api/items/{item_id}/ai-yield")
    s = SessionLocal()
    rows = s.query(ExpertYield).filter(ExpertYield.item_id == item_id).all()
    assert len(rows) == 1 and rows[0].good_percent == 55
    s.close()


def test_ai_yield_honest_when_no_data(catalog_client, monkeypatch):
    from app import ai_yield
    from app.db.models import Item
    from app.db.session import SessionLocal
    from app.normalize import normalize_name

    monkeypatch.setenv("PERPLEXITY_API_KEY", "test-key")
    monkeypatch.setattr(ai_yield, "ask_yield", lambda name, family: {
        "good_percent": None, "sources": [], "comment": "нет данных"})
    s = SessionLocal()
    it = Item(name="Блок ТМСП-3 б/у", unit="шт", family="блок ТМС",
              name_normalized=normalize_name("Блок ТМСП-3 б/у"), is_trade=True)
    s.add(it)
    s.commit()
    item_id = it.id
    s.close()
    r = catalog_client.post(f"/api/items/{item_id}/ai-yield")
    assert r.status_code == 200
    assert r.json() == {"saved": False, "comment": "нет данных"}


# ------------------------------------------------------ сводка и статусы КП


def test_kp_list_uses_cached_totals(client):
    """Список КП не пересчитывает позиции: итоги берутся из кэша расчёта."""
    docs = client.get("/api/kp").json()
    assert docs[0]["base_buyout"] is None  # КП ещё не открывали
    client.get("/api/kp/1")                # открытие обновляет кэш
    docs = client.get("/api/kp").json()
    assert docs[0]["base_buyout"] is not None
    assert docs[0]["calc_at"] and docs[0]["weight_t"] == pytest.approx(10.0)


def test_kp_status_transitions_stamp_dates(client):
    r = client.patch("/api/kp/1", json={"status": "отправлено",
                                        "offer_amount": 1_000_000})
    assert r.status_code == 200 and r.json()["status"] == "отправлено"
    doc = next(d for d in client.get("/api/kp").json() if d["id"] == 1)
    assert doc["sent_at"] is not None          # дата проставлена сама
    assert doc["offer_amount"] == 1_000_000

    client.patch("/api/kp/1", json={"status": "выиграно"})
    assert client.get("/api/kp/1").json()["status"] == "выиграно"
    assert client.patch("/api/kp/1", json={"status": "нет-такого"}).status_code == 422


def test_kp_funnel_counts_and_win_rate(client):
    from app.db.models import KpDocument
    from app.db.session import SessionLocal

    s = SessionLocal()
    for status, amount in (("отправлено", 500_000), ("выиграно", 300_000),
                           ("проиграно", 200_000)):
        s.add(KpDocument(title=f"КП {status}", status=status,
                         offer_amount=amount, calc_weight_t=10))
    s.commit()
    s.close()

    summary = client.get("/api/kp-summary").json()
    stages = {st["status"]: st for st in summary["stages"]}
    assert stages["отправлено"]["count"] == 1
    assert stages["выиграно"]["offer_sum"] == pytest.approx(300_000)
    assert stages["проиграно"]["count"] == 1
    # доля побед считается только от решённых (1 из 2)
    assert summary["win_rate"] == pytest.approx(50.0)
    assert summary["total"] == 4


# ------------------------------------------------ массовое слияние дублей


def test_merge_tier_classification():
    """Уровень очевидности: пунктуация — безопасно, «б/у» — на решение."""
    from types import SimpleNamespace as NS
    from app.importers.nomenclature import merge_tier

    def item(name, unit="т"):
        return NS(name=name, unit=unit)

    assert merge_tier(item("*Металлолом 5А"), item("Металлолом 5А")) == "safe"
    assert merge_tier(item("Труба НКТ 73*5,5 б/у*"),
                      item("Труба НКТ 73х5.5 б/у")) == "safe"
    assert merge_tier(item("Труба 325*8"), item("Труба 325*8 б/у")) == "bu"
    # разные семейства единиц — не дубли ни при каких условиях
    assert merge_tier(item("Труба НКТ 89 б/у", "т"),
                      item("Труба б/у НКТ 89", "м")) is None
    assert merge_tier(item("Труба 2 1/2 б/у"), item("Труба 3 б/у")) is None


def test_auto_merge_preview_then_apply(catalog_client):
    """Предпросмотр не меняет данные; применение сливает и сохраняет guid."""
    from app.db.models import Item, ItemAlias, MergeCandidate, PriceQuote
    from app.db.session import SessionLocal
    from app.normalize import normalize_name

    s = SessionLocal()
    a = Item(name="*Лом 12А", name_normalized=normalize_name("*Лом 12А"),
             unit="т", family="лом", is_trade=True, guid="guid-a")
    b = Item(name="Лом 12А", name_normalized=normalize_name("Лом 12А"),
             unit="т", family="лом", is_trade=True, guid="guid-b")
    s.add_all([a, b])
    s.commit()
    # у «чистой» карточки есть история — она и должна выжить
    s.add(PriceQuote(item_id=b.id, quote_type="sales_fact", price=20_000,
                     unit="т", ttl_days=9999))
    s.add(MergeCandidate(item_id_a=a.id, item_id_b=b.id, score=1.0))
    s.commit()
    a_id, b_id = a.id, b.id
    s.close()

    preview = catalog_client.post("/api/merge-candidates/auto",
                                  json={"tier": "safe"}).json()
    assert preview["found"] >= 1 and preview["merged"] == 0
    assert any(p["keep"] == "Лом 12А" for p in preview["preview"])

    applied = catalog_client.post("/api/merge-candidates/auto",
                                  json={"tier": "safe", "apply": True}).json()
    assert applied["merged"] >= 1 and not applied["failed"]

    s = SessionLocal()
    keep, drop = s.get(Item, b_id), s.get(Item, a_id)
    assert keep.is_trade and not drop.is_trade   # выжила карточка с историей
    aliases = {al.guid for al in s.query(ItemAlias).filter(
        ItemAlias.item_id == b_id)}
    assert "guid-a" in aliases                   # guid из 1С сохранён
    s.close()


def test_merge_refuses_incompatible_units(catalog_client):
    """Штучную и весовую карточки слить нельзя — это разные сущности."""
    from app.db.models import Item, MergeCandidate
    from app.db.session import SessionLocal
    from app.normalize import normalize_name

    s = SessionLocal()
    a = Item(name="Труба НКТ 89 б/у (тн)", unit="т", is_trade=True,
             name_normalized=normalize_name("Труба НКТ 89 б/у (тн)"))
    b = Item(name="Труба б/у НКТ 89", unit="м", is_trade=True,
             name_normalized=normalize_name("Труба б/у НКТ 89"))
    s.add_all([a, b])
    s.commit()
    cand = MergeCandidate(item_id_a=a.id, item_id_b=b.id, score=1.0)
    s.add(cand)
    s.commit()
    cand_id, a_id, b_id = cand.id, a.id, b.id
    s.close()

    # в автослияние такая пара не попадает ни на одном уровне
    for tier in ("safe", "bu"):
        res = catalog_client.post("/api/merge-candidates/auto",
                                  json={"tier": tier, "apply": True}).json()
        assert all(p["keep"] != "Труба НКТ 89 б/у (тн)"
                   for p in res["preview"]), tier
    # ручное слияние тоже отбивается
    r = catalog_client.post(f"/api/merge-candidates/{cand_id}/merge")
    assert r.status_code == 422 and "единиц" in r.json()["detail"]
    s = SessionLocal()
    assert s.get(Item, a_id).is_trade and s.get(Item, b_id).is_trade
    s.close()


def test_merge_refuses_chain(catalog_client):
    """Уже поглощённую карточку нельзя слить повторно: иначе один и тот же
    alias оказывается у двух владельцев (нашли на реальном справочнике)."""
    from app.db.models import Item
    from app.db.session import SessionLocal
    from app.importers.nomenclature import merge_items
    from app.normalize import normalize_name

    s = SessionLocal()
    a, b, c = (Item(name=n, name_normalized=normalize_name(n), unit="т",
                    is_trade=True, guid=f"g-{n}")
               for n in ("Труба 100 б/у", "Труба 100 б/у*", "Труба 100 б/у**"))
    s.add_all([a, b, c])
    s.commit()
    merge_items(s, a, c, "тест")     # c поглощена карточкой a
    s.commit()
    with pytest.raises(ValueError, match="уже объединена"):
        merge_items(s, b, c, "тест")  # повторное слияние той же c
    with pytest.raises(ValueError, match="сам"):
        merge_items(s, a, a, "тест")
    s.close()


# ---------------------------------------------------- ИИ-оценка массы


def test_ai_weight_refused_for_weight_units(catalog_client, monkeypatch):
    """Для позиций в тоннах масса единицы задана самой единицей."""
    from app.db.models import Item
    from app.db.session import SessionLocal
    from app.normalize import normalize_name

    monkeypatch.setenv("PERPLEXITY_API_KEY", "test-key")
    s = SessionLocal()
    it = Item(name="Труба 89 б/у (т)", unit="т", is_trade=True, family="труба",
              name_normalized=normalize_name("Труба 89 б/у (т)"))
    s.add(it)
    s.commit()
    item_id = it.id
    s.close()
    r = catalog_client.post(f"/api/items/{item_id}/ai-weight")
    assert r.status_code == 422 and "1 т = 1000 кг" in r.json()["detail"]


def test_ai_weight_rejects_implausible_mass(catalog_client, monkeypatch):
    """Масса вне разумного коридора — ошибка размерности, не сохраняем."""
    from app import ai_weight
    from app.db.models import Item
    from app.db.session import SessionLocal
    from app.normalize import normalize_name

    monkeypatch.setenv("PERPLEXITY_API_KEY", "test-key")
    s = SessionLocal()
    it = Item(name="Двигатель ПЭД-45 б/у", unit="шт", is_trade=True, family="ПЭД",
              name_normalized=normalize_name("Двигатель ПЭД-45 б/у"))
    s.add(it)
    s.commit()
    item_id = it.id
    s.close()

    monkeypatch.setattr(ai_weight, "ask_weight", lambda n, u: {
        "mass_kg": 850_000, "sources": [], "comment": "масса партии"})
    r = catalog_client.post(f"/api/items/{item_id}/ai-weight")
    assert r.status_code == 200 and r.json()["saved"] is False
    assert "размерности" in r.json()["comment"]

    monkeypatch.setattr(ai_weight, "ask_weight", lambda n, u: {
        "mass_kg": 780, "confidence": "medium",
        "sources": ["https://example.com/ped"], "comment": "паспорт завода"})
    r = catalog_client.post(f"/api/items/{item_id}/ai-weight")
    assert r.json()["saved"] and r.json()["mass_kg"] == 780
    card = catalog_client.get(f"/api/items/{item_id}").json()
    assert card["unit_mass_kg"] == 780
    assert card["unit_mass_url"] == "https://example.com/ped"
    assert card["unit_mass_confidence"] == "medium"


def test_manual_weight_marked_high_confidence(catalog_client):
    from app.db.models import Item
    from app.db.session import SessionLocal
    from app.normalize import normalize_name

    s = SessionLocal()
    it = Item(name="Гидрозащита ПР92 б/у", unit="шт", is_trade=True,
              family="гидрозащита",
              name_normalized=normalize_name("Гидрозащита ПР92 б/у"))
    s.add(it)
    s.commit()
    item_id = it.id
    s.close()
    r = catalog_client.patch(f"/api/items/{item_id}", json={"unit_mass_kg": 95})
    assert r.status_code == 200
    card = catalog_client.get(f"/api/items/{item_id}").json()
    assert card["unit_mass_kg"] == 95
    assert card["unit_mass_source"] == "введено вручную"
    assert card["unit_mass_confidence"] == "high"


def test_ai_weight_rejects_answer_contradicting_our_disassembly(
        catalog_client, monkeypatch):
    """Реальный случай: секцию ЭЦН модель приняла за кабель и выдала 2 кг
    вместо ~270, причём с доверием medium. Сверка с профилем это отбивает."""
    from app import ai_weight
    from app.db.models import ExpertMetalProfile, Item
    from app.db.session import SessionLocal
    from app.normalize import normalize_name

    monkeypatch.setenv("PERPLEXITY_API_KEY", "test-key")
    s = SessionLocal()
    s.add(ExpertMetalProfile(family="секция ЭЦН", material="чермет",
                             kg_per_unit=271))
    it = Item(name="Секция 117ЭЦН(НГ) 5-25 (4м) б/у", unit="шт", is_trade=True,
              family="секция ЭЦН",
              name_normalized=normalize_name("Секция 117ЭЦН(НГ) 5-25 (4м) б/у"))
    s.add(it)
    s.commit()
    item_id = it.id
    s.close()

    monkeypatch.setattr(ai_weight, "ask_weight", lambda n, u: {
        "mass_kg": 2, "confidence": "medium", "sources": [],
        "comment": "ВВГнг(А)-LS 5х25, 1,95 кг/м × 4 м"})
    body = catalog_client.post(f"/api/items/{item_id}/ai-weight").json()
    assert body["saved"] is False
    assert "расходится с нашими разборками" in body["comment"]
    assert body["profile_kg"] == 271
    assert catalog_client.get(f"/api/items/{item_id}").json()["unit_mass_kg"] is None

    # правдоподобный ответ в пределах допуска сохраняется
    monkeypatch.setattr(ai_weight, "ask_weight", lambda n, u: {
        "mass_kg": 300, "confidence": "high", "sources": [], "comment": "каталог"})
    assert catalog_client.post(f"/api/items/{item_id}/ai-weight").json()["saved"]


def test_auth_off_allows_anonymous_and_names_the_mode(client):
    """METOPTTORG_AUTH=off: вход не требуется, в журнале так и пишется."""
    assert client.get("/api/kp").status_code == 200      # без пароля
    me = client.get("/api/me").json()
    assert me["login"] == "без входа"
    assert me["role"] == "директор" and me["can_approve"] is True
    # правка нормативов доступна — роль директорская
    assert client.post("/api/norms",
                       json={"key": "target_margin_percent",
                             "value": 27}).status_code == 200


def test_header_links_to_portal(client, monkeypatch):
    """Логотип в шапке ведёт на портал: адрес из настройки службы."""
    assert client.get("/api/config").json()["portal_url"] is None
    monkeypatch.setenv("PORTAL_URL", "http://192.168.6.157:8079/")
    assert client.get("/api/config").json()["portal_url"] == \
        "http://192.168.6.157:8079/"
    html = client.get("/").text
    assert 'id="homeLink"' in html and "портала" in html


# ------------------------------------------------------------ рынок лома


def _make_review(path):
    """Мини-обзор в формате реального файла: содержание + два листа цен."""
    import datetime as dt
    import openpyxl

    wb = openpyxl.Workbook()
    c = wb.active
    c.title = "Содержание"
    c["E1"] = "выпуск 24.08.2026"
    c["E3"] = "Неделя № 34: 17.08.2026 - 23.08.2026"
    days = [dt.datetime(2026, 8, d) for d in range(17, 24)]
    for sheet, rows in (
            ("Цены по обл", [("Пермский кр.", 27000, 27100, 27200),
                             ("респ. Коми", 21500, 21700, 22000)]),
            ("Цены по потр", [("Магнитогорский МК", 30300, 30500, 30700)])):
        ws = wb.create_sheet(sheet)
        ws["A1"] = "Средневзвешенная цена"
        for i, d in enumerate(days[:3], start=2):
            ws.cell(row=2, column=i, value=d)
        for r, (name, *vals) in enumerate(rows, start=3):
            ws.cell(row=r, column=1, value=name)
            for i, v in enumerate(vals, start=2):
                ws.cell(row=r, column=i, value=v)
    wb.save(path)


def test_market_review_import_and_period(client, tmp_path):
    """Период берётся из строки недели, а не из даты выпуска."""
    f = tmp_path / "review.xlsx"
    _make_review(f)
    with open(f, "rb") as fh:
        r = client.post("/api/market-review/upload",
                        files={"file": ("review.xlsx", fh)})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["issue"] == "34-2026"
    assert body["regions"] == 2 and body["consumers"] == 1
    assert body["period"][0].startswith("2026-08-17")   # начало недели
    assert body["period"][1].startswith("2026-08-23")   # а не выпуск 24.08

    prices = client.get("/api/market-prices?kind=region").json()
    perm = next(p for p in prices["prices"] if p["scope"] == "Пермский кр.")
    assert perm["price"] == pytest.approx(27100)        # среднее по дням
    assert perm["basis"] == "FCA" and perm["days"] == 3
    cons = client.get("/api/market-prices?kind=consumer").json()
    assert cons["prices"][0]["basis"] == "CPT"


def test_market_ratio_only_for_ferrous(client, tmp_path):
    """Индекс — про чёрный лом: коэффициент для цветмета запрещён."""
    from app.calc.market import INDEXED_MATERIALS, set_ratio
    from app.db.session import SessionLocal

    assert INDEXED_MATERIALS == {"чермет"}
    s = SessionLocal()
    with pytest.raises(ValueError, match="не индексируется"):
        set_ratio(s, "медь", 5.17, "ошибочная привязка")
    s.close()
    r = client.post("/api/scrap-prices/ratio",
                    json={"material": "медь", "ratio": 5.17})
    assert r.status_code in (422, 500)


def test_recommended_price_is_index_times_ratio(client, tmp_path):
    """Рекомендация по чермету = индекс × наш коэффициент, с пояснением."""
    from app.calc.market import set_ratio
    from app.db.session import SessionLocal

    f = tmp_path / "review2.xlsx"
    _make_review(f)
    with open(f, "rb") as fh:
        client.post("/api/market-review/upload",
                    files={"file": ("review2.xlsx", fh)})
    s = SessionLocal()
    set_ratio(s, "чермет", 0.68, "медиана наших продаж 18 400 ₽/т")
    s.close()

    rec = client.get("/api/scrap-prices/recommended").json()["чермет"]
    assert rec["index"] == pytest.approx(27100)
    assert rec["recommended"] == pytest.approx(27100 * 0.68, rel=1e-3)
    assert "×" in rec["explanation"] and "FCA" in rec["explanation"]


# ------------------------------------------------------ вкладка «Клиенты»


def test_leads_campaign_runs_without_ai_key(client, monkeypatch):
    """Без ключа провайдера подбор всё равно доходит до конца и объясняет почему.

    Внутренние лиды и ценовой коридор считаются по фактам 1С, интернет нужен
    только для внешних каналов — иначе вкладка была бы бесполезна на сервере,
    где провайдер закрыт периметром.
    """
    monkeypatch.delenv("OPENROUTER_API_KEY", raising=False)
    monkeypatch.delenv("PERPLEXITY_API_KEY", raising=False)
    from app.db.models import Counterparty, DealFact, Item, PriceQuote
    from app.db.session import SessionLocal
    from app.leads import service as leads_service
    from app.normalize import normalize_name
    import datetime as dt

    s = SessionLocal()
    # Как в настоящей базе: цена продажи лежит на карточке-двойнике, а на той,
    # с которой приходит КП, её нет.
    twin_name = "Погружной электродвигатель 10ВДМ100-2400-3,0-117 (шт)"
    it = Item(name="ВДМ 100-2400-3.0-117В5 б/у", unit="шт", is_trade=True,
              family="ПЭД", size_key="вдм-100",
              name_normalized=normalize_name("ВДМ 100-2400-3.0-117В5 б/у"))
    twin = Item(name=twin_name, unit="шт", is_trade=True, family="ПЭД",
                size_key="вдм-100", name_normalized=normalize_name(twin_name))
    cp = Counterparty(name="СОЮЗ-ТЕХНО ООО",
                      name_normalized=normalize_name("СОЮЗ-ТЕХНО ООО"))
    s.add_all([it, twin, cp])
    s.commit()
    s.add(PriceQuote(item_id=twin.id, quote_type="sales_fact", price=35_625,
                     unit="шт", ttl_days=9999))
    s.commit()
    s.add(DealFact(counterparty_id=cp.id, item_name="10ВДМ100-2400-3,0-117 (шт)",
                   family="ПЭД", size_key="вдм-100", series_mark="вдм",
                   direction="sale", period=dt.datetime(2025, 9, 1), quantity=2,
                   unit="шт", price=35_625, revenue=71_250, is_scrap=False))
    s.commit()
    item_id = it.id
    s.close()

    # кампания идёт в фоне; в тесте фон отключаем и гоняем её синхронно,
    # иначе проверка гонялась бы с потоком
    monkeypatch.setattr(leads_service, "start", lambda *a, **k: None)
    r = client.post("/api/leads/search",
                    json={"item_id": item_id, "unit": "шт", "quantity": 3})
    assert r.status_code == 200, r.text
    campaign_id = r.json()["campaign_id"]

    s = SessionLocal()
    leads_service.run(s, campaign_id)
    s.close()

    d = client.get(f"/api/leads/campaigns/{campaign_id}").json()
    assert d["status"] == "готово"
    assert "ключ" in (d["error"] or ""), "должно быть сказано, почему нет внешних"
    assert d["price"]["target"] == 35_625
    # Цель вкладки — новые покупатели. Наш действующий клиент в основной
    # список лидов не попадает: он идёт отдельно, как ценовой ориентир.
    assert [l["name"] for l in d["existing_customers"]] == ["СОЮЗ-ТЕХНО ООО"]
    assert all(not l["is_existing"] for l in d["leads"])

    lead_id = d["existing_customers"][0]["id"]
    assert client.patch(f"/api/leads/{lead_id}",
                        json={"status": "в работе"}).status_code == 200
    assert client.patch(f"/api/leads/{lead_id}",
                        json={"status": "куплено"}).status_code == 422
    draft = client.get(f"/api/leads/{lead_id}/draft").json()
    assert "МетОптТорг" in draft["body"]
    assert "не отправляет" in draft["warning"]
    assert client.get(f"/api/leads/campaigns/{campaign_id}/export").status_code == 200


def test_price_ladder_endpoint_explains_the_number(client):
    """«Откуда цифра» — эндпоинт возвращает все сработавшие ступени каскада."""
    r = client.get("/api/leads/price-ladder",
                   params={"name": "Труба НКТ 73 б/у", "unit": "т"})
    assert r.status_code == 200
    assert "candidates" in r.json()
