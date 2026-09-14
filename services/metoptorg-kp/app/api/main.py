"""FastAPI-приложение: КП, справочник, прайс лома, нормативы, экспорт Excel."""
from __future__ import annotations

import hashlib
import shutil
from pathlib import Path

from fastapi import Depends, FastAPI, File, HTTPException, UploadFile
from fastapi.responses import FileResponse, HTMLResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session

from ..calc.engine import calc_position
from ..db.models import (
    ApprovedValue,
    Item,
    KpDocument,
    KpPosition,
    ScrapPrice,
    SourceFile,
)
from ..db.session import get_session, init_db
from ..importers.cable_files import (
    import_brand_table,
    import_valuation_sheet,
    import_vnpz_sheet,
)
from ..importers.kp import import_kp, match_item
from ..importers.metal_yields import (
    XlsxScanError,
    import_yields_file,
    rebuild_profiles,
)
from ..importers.nomenclature import import_nomenclature
from ..seed import seed
from ..db.session import SessionLocal

import datetime as _dt

from ..auth import User, current_user, log_action, require_director


def dt_now():
    return _dt.datetime.now(_dt.timezone.utc)

# аутентификация обязательна на всех маршрутах, включая интерфейс
app = FastAPI(title="МетОптТорг — оценка КП",
              dependencies=[Depends(current_user)])

from ..db.session import DATA_DIR

BASE_DIR = Path(__file__).resolve().parents[2]
UPLOADS = DATA_DIR / "uploads"
EXPORTS = DATA_DIR / "exports"
WEB = BASE_DIR / "web"


@app.on_event("startup")
def _startup():
    init_db()
    # каталоги данных должны существовать до первого запроса: экспорт КП
    # пишет файл сюда же, а загрузки может ещё не быть ни одной
    UPLOADS.mkdir(parents=True, exist_ok=True)
    EXPORTS.mkdir(parents=True, exist_ok=True)
    s = SessionLocal()
    try:
        seed(s)
    finally:
        s.close()
    # фоновый опрос каналов; MONITOR_INTERVAL_MIN=0 выключает
    from ..monitor.service import start_scheduler
    start_scheduler(SessionLocal)
    # кампании подбора живут в потоках и не переживают рестарт службы —
    # оборванные помечаем, чтобы оператор не ждал их вечно
    from ..leads.service import release_orphans
    s = SessionLocal()
    try:
        release_orphans(s)
    finally:
        s.close()


# ---------------------------------------------------------------- загрузка


def _store_file(session: Session, up: UploadFile, kind: str) -> SourceFile:
    UPLOADS.mkdir(parents=True, exist_ok=True)
    tmp = UPLOADS / f"tmp_{up.filename}"
    with tmp.open("wb") as f:
        shutil.copyfileobj(up.file, f)
    sha = hashlib.sha256(tmp.read_bytes()).hexdigest()
    existing = session.query(SourceFile).filter(SourceFile.sha256 == sha).first()
    if existing:
        tmp.unlink()
        raise HTTPException(
            409, f"Этот файл уже загружен как «{existing.filename}» "
                 f"({existing.uploaded_at:%d.%m.%Y}) — дубликат по SHA-256.")
    dest = UPLOADS / f"{sha[:12]}_{up.filename}"
    tmp.rename(dest)
    sf = SourceFile(filename=up.filename, sha256=sha, kind=kind,
                    stored_path=str(dest))
    session.add(sf)
    session.commit()
    return sf


@app.post("/api/kp/upload")
def upload_kp(file: UploadFile = File(...), session: Session = Depends(get_session)):
    sf = _store_file(session, file, "kp")
    try:
        doc, res = import_kp(session, sf.stored_path, sf.filename, sf.id)
        sf.parse_status = "parsed"
        session.commit()
    except ValueError as e:
        sf.parse_status = "error"
        sf.parse_log = str(e)
        session.commit()
        raise HTTPException(422, str(e))
    return {"document_id": doc.id, "positions": res.positions, "lots": res.lots,
            "sheet": res.sheet, "total_weight_t": round(res.total_weight_t, 3),
            "file_total_t": (round(res.file_total_t, 3) if res.file_total_t
                             else None),
            "warnings": res.warnings}


@app.post("/api/yields/upload")
def upload_yields(file: UploadFile = File(...),
                  session: Session = Depends(get_session)):
    sf = _store_file(session, file, "yields")
    try:
        st = import_yields_file(session, sf.stored_path, sf.filename, sf.id)
        sf.parse_status = "parsed"
        session.commit()
        profiles = rebuild_profiles(session)
        return {"batches": st.batches, "yields": st.yields, "lists": st.lists,
                "profiles_rebuilt": profiles}
    except XlsxScanError as e:
        sf.parse_status = "needs_review"
        sf.parse_log = str(e)
        session.commit()
        raise HTTPException(422, str(e))


@app.post("/api/nomenclature/upload")
def upload_nomenclature(file: UploadFile = File(...),
                        session: Session = Depends(get_session),
                        user: User = Depends(require_director)):
    sf = _store_file(session, file, "nomenclature")
    st = import_nomenclature(session, sf.stored_path)
    sf.parse_status = "parsed"
    session.commit()
    return {"rows": st.rows, "created": st.created,
            "merged_duplicates": st.merged_duplicates, "trade": st.trade,
            "fuzzy_candidates": st.fuzzy_candidates}


# ---------------------------------------------------------------- КП


def _position_payload(session: Session, p: KpPosition,
                      as_of=None) -> dict:
    from ..db.models import PriceQuote
    calc = calc_position(session, p, as_of=as_of)
    item = session.get(Item, p.item_id) if p.item_id else None
    benchmarks = {}
    if item:
        for qt in ("purchase_fact", "sales_fact"):
            row = (session.query(PriceQuote)
                   .filter(PriceQuote.item_id == item.id,
                           PriceQuote.quote_type == qt)
                   .order_by(PriceQuote.quoted_at.desc()).first())
            if row:
                benchmarks[qt] = {"price": row.price, "unit": row.unit,
                                  "notes": row.notes, "source": row.source}

    def scen(r):
        return {
            "good_percent": round(r.good_percent, 1),
            "good_source": r.good_source,
            "scrap_value": round(r.scrap_value, 2),
            "scrap_basis": r.scrap_basis,
            "good_value": round(r.good_value, 2),
            "good_price_source": r.good_price_source,
            "good_floor": r.good_floor_applied,
            "resale_total": round(r.resale_total, 2),
            "buyout_total": round(r.buyout_total, 2),
            "buyout_per_unit": round(r.buyout_per_unit, 2) if r.buyout_per_unit else None,
            "unit_mass_kg": round(r.unit_mass_kg, 2) if r.unit_mass_kg else None,
            "mass_source": r.mass_source,
            "warnings": r.warnings,
        }

    return {
        "id": p.id,
        "raw_name": p.raw_name,
        "quantity": p.quantity,
        "unit": p.unit,
        "weight_kg": p.weight_kg,
        "calc_weight_kg": (round(calc.total_weight_kg, 1)
                           if calc.total_weight_kg else None),
        "lots_merged": p.lots_merged,
        "comment": p.comment,
        "item": {"id": item.id, "name": item.name, "unit": item.unit} if item else None,
        "match_kind": p.match_kind,
        "match_score": p.match_score,
        "manual": {
            "good_percent": p.manual_good_percent,
            "resale_price": p.manual_resale_price,
            "scrap_price": p.manual_scrap_price,
            "quantity": p.manual_quantity,
            "weight_kg": p.manual_weight_kg,
        },
        "base": scen(calc.base),
        "bp": scen(calc.bp),
        "cable_value_per_tonne": (round(calc.cable_value_per_tonne)
                                  if calc.cable_value_per_tonne else None),
        "benchmarks": benchmarks,
    }


def _norm_value_for(session: Session, key: str, scope: str | None) -> dict:
    """Действующая ставка и откуда она взята — для показа в интерфейсе."""
    from ..calc.engine import _norm_value
    value = _norm_value(session, key, 0.0, scope=scope)
    scoped = None
    if scope:
        scoped = (session.query(ApprovedValue)
                  .filter(ApprovedValue.key == key,
                          ApprovedValue.is_current.is_(True),
                          ApprovedValue.scope == scope).first())
    return {"value": value,
            "source": (f"факт по месту «{scope}»" if scoped
                       else "общая медиана по компании")}


@app.get("/api/kp")
def list_kp(session: Session = Depends(get_session)):
    """Список КП с сохранёнными итогами: пересчёт сотен позиций тут не нужен."""
    docs = session.query(KpDocument).order_by(KpDocument.created_at.desc()).all()
    return [{"id": d.id, "title": d.title, "seller": d.seller,
             "seller_region": d.seller_region, "status": d.status or "черновик",
             "created_at": d.created_at.isoformat() if d.created_at else None,
             "sent_at": d.sent_at.isoformat() if d.sent_at else None,
             "positions": len(d.positions),
             "weight_t": d.calc_weight_t,
             "base_buyout": d.calc_base_buyout, "bp_buyout": d.calc_bp_buyout,
             "offer_amount": d.offer_amount, "offer_scenario": d.offer_scenario,
             "comment": d.comment,
             "calc_at": d.calc_at.isoformat() if d.calc_at else None}
            for d in docs]


@app.get("/api/kp-summary")
def kp_summary(session: Session = Depends(get_session)):
    """Воронка: сколько КП и денег на каждой стадии."""
    stages = ["черновик", "отправлено", "выиграно", "проиграно"]
    docs = session.query(KpDocument).all()
    out = {st: {"count": 0, "offer_sum": 0.0, "weight_t": 0.0} for st in stages}
    for d in docs:
        st = d.status if d.status in out else "черновик"
        out[st]["count"] += 1
        # для отправленных берём то, что реально предложили, иначе расчёт
        out[st]["offer_sum"] += (d.offer_amount
                                 or d.calc_base_buyout or 0.0)
        out[st]["weight_t"] += d.calc_weight_t or 0.0
    won = out["выиграно"]["count"]
    decided = won + out["проиграно"]["count"]
    return {"stages": [{"status": st, **out[st]} for st in stages],
            "win_rate": round(won / decided * 100, 1) if decided else None,
            "total": len(docs)}


@app.get("/api/kp/{doc_id}")
def get_kp(doc_id: int, session: Session = Depends(get_session)):
    doc = session.get(KpDocument, doc_id)
    if not doc:
        raise HTTPException(404, "КП не найдено")
    positions = [_position_payload(session, p, as_of=doc.priced_at)
                 for p in doc.positions]
    totals = {
        "base_resale": round(sum(p["base"]["resale_total"] for p in positions), 2),
        "bp_resale": round(sum(p["bp"]["resale_total"] for p in positions), 2),
        "base_buyout": round(sum(p["base"]["buyout_total"] for p in positions), 2),
        "bp_buyout": round(sum(p["bp"]["buyout_total"] for p in positions), 2),
        "weight_t": round(sum((p["weight_kg"] or p["calc_weight_kg"] or 0)
                              for p in positions) / 1000.0, 3),
    }
    imported_t = totals["weight_t"]
    reconciliation = None
    if doc.file_total_weight_kg:
        file_t = doc.file_total_weight_kg / 1000.0
        diff = imported_t - file_t
        reconciliation = {
            "file_total_t": round(file_t, 3),
            "imported_t": round(imported_t, 3),
            "diff_t": round(diff, 3),
            "diff_percent": round(diff / file_t * 100, 2) if file_t else None,
            "ok": abs(diff) / file_t <= 0.005 if file_t else False,
        }
    # сохраняем итоги, чтобы список КП и воронка не пересчитывали позиции
    doc.calc_base_buyout = totals["base_buyout"]
    doc.calc_bp_buyout = totals["bp_buyout"]
    doc.calc_weight_t = totals["weight_t"]
    doc.calc_at = dt_now()
    session.commit()
    logistics = _norm_value_for(session, "logistics_cost_per_tonne",
                                doc.seller_region)
    return {"id": doc.id, "title": doc.title, "positions": positions,
            "totals": totals, "reconciliation": reconciliation,
            "seller": doc.seller, "seller_region": doc.seller_region,
            "status": doc.status or "черновик", "offer_amount": doc.offer_amount,
            "offer_scenario": doc.offer_scenario, "comment": doc.comment,
            "logistics_rate": logistics,
            "priced_at": doc.priced_at.isoformat() if doc.priced_at else None}


class ManualPatch(BaseModel):
    good_percent: float | None = None
    resale_price: float | None = None
    scrap_price: float | None = None
    quantity: float | None = None
    weight_kg: float | None = None
    comment: str | None = None
    item_id: int | None = None


@app.patch("/api/kp/position/{pos_id}")
def patch_position(pos_id: int, patch: ManualPatch,
                   session: Session = Depends(get_session),
                   user: User = Depends(current_user)):
    p = session.get(KpPosition, pos_id)
    if not p:
        raise HTTPException(404, "Позиция не найдена")
    fields = patch.model_dump(exclude_unset=True)
    mapping = {"good_percent": "manual_good_percent",
               "resale_price": "manual_resale_price",
               "scrap_price": "manual_scrap_price",
               "quantity": "manual_quantity",
               "weight_kg": "manual_weight_kg",
               "comment": "comment"}
    for k, v in fields.items():
        if k == "item_id":
            p.item_id = v
            p.match_kind = "manual"
            p.match_score = 1.0
        elif k in mapping:
            setattr(p, mapping[k], v)
    p.updated_by = user.login
    p.updated_at = dt_now()
    session.commit()
    log_action(session, user, "корректировка позиции", "kp_positions", p.id,
               "; ".join(f"{k}={v}" for k, v in fields.items()))
    return _position_payload(session, p)  # мгновенный пересчёт


@app.get("/api/items/search")
def search_items(q: str, session: Session = Depends(get_session)):
    from ..normalize import normalize_name
    norm = normalize_name(q)
    rows = (session.query(Item)
            .filter(Item.is_trade.is_(True), Item.name_normalized.contains(norm))
            .order_by(Item.name_normalized).limit(30).all())
    return [{"id": i.id, "name": i.name, "unit": i.unit, "family": i.family}
            for i in rows]


# ---------------------------------------------------------------- прайс/нормативы


@app.get("/api/scrap-prices")
def scrap_prices(session: Session = Depends(get_session)):
    rows = (session.query(ScrapPrice).filter(ScrapPrice.is_current.is_(True))
            .order_by(ScrapPrice.material).all())
    return [{"id": r.id, "material": r.material, "grade": r.grade,
             "price_per_tonne": r.price_per_tonne, "source": r.source,
             "valid_from": r.valid_from.isoformat() if r.valid_from else None}
            for r in rows]


class ScrapPriceIn(BaseModel):
    material: str
    grade: str | None = None
    price_per_tonne: float
    source: str = "ручной ввод"


@app.post("/api/scrap-prices")
def set_scrap_price(body: ScrapPriceIn, session: Session = Depends(get_session),
                    user: User = Depends(require_director)):
    # история: старую запись закрываем, новую открываем
    session.query(ScrapPrice).filter(
        ScrapPrice.material == body.material,
        ScrapPrice.is_current.is_(True)).update({"is_current": False})
    sp = ScrapPrice(material=body.material, grade=body.grade,
                    price_per_tonne=body.price_per_tonne, source=body.source,
                    created_by=user.login)
    session.add(sp)
    session.commit()
    return {"ok": True, "id": sp.id}


@app.post("/api/kp/position/{pos_id}/ai-price")
def ai_price(pos_id: int, session: Session = Depends(get_session)):
    """ИИ-поиск цены б/у для позиции (нужен ключ провайдера в окружении)."""
    from .. import ai_price as ai
    p = session.get(KpPosition, pos_id)
    if not p:
        raise HTTPException(404, "Позиция не найдена")
    if not p.item_id:
        raise HTTPException(422, "Позиция не сопоставлена со справочником")
    try:
        result = ai.ai_price_for_item(session, p.item_id)
    except ai.AiUnavailable as e:
        raise HTTPException(503, str(e))
    except (ValueError, KeyError) as e:
        raise HTTPException(502, f"не удалось разобрать ответ ИИ: {e}")
    return {"result": result, "position": _position_payload(session, p)}


@app.post("/api/items/{item_id}/ai-price")
def ai_price_item(item_id: int, session: Session = Depends(get_session)):
    """ИИ-поиск цены для карточки справочника (вне контекста КП)."""
    from .. import ai_price as ai
    try:
        return ai.ai_price_for_item(session, item_id)
    except ai.AiUnavailable as e:
        raise HTTPException(503, str(e))
    except ValueError as e:
        raise HTTPException(422, str(e))


@app.get("/api/ai-status")
def ai_status():
    from .. import ai_price as ai
    return {"enabled": ai.is_enabled()}


@app.get("/api/expert-yields")
def list_expert_yields(session: Session = Depends(get_session)):
    from ..db.models import ExpertYield
    rows = (session.query(ExpertYield)
            .order_by(ExpertYield.block, ExpertYield.category,
                      ExpertYield.size_key).all())
    return [{"id": r.id, "block": r.block, "category": r.category,
             "size_key": r.size_key, "item_name": r.item_name,
             "good_percent": r.good_percent, "expert_name": r.expert_name,
             "confidence": r.confidence, "notes": r.notes} for r in rows]


class ExpertYieldIn(BaseModel):
    block: str  # mot | expert | ai
    good_percent: float
    category: str | None = None  # семейство: ПЭД, секция ЭЦН, гидрозащита…
    size_key: str | None = None  # нкт-73, труба-159…
    item_name: str | None = None  # конкретная позиция (бьёт размер и категорию)
    expert_name: str = ""
    notes: str | None = None


@app.post("/api/expert-yields")
def approve_expert_yield(body: ExpertYieldIn,
                         session: Session = Depends(get_session),
                         user: User = Depends(require_director)):
    from ..db.models import ExpertYield
    import datetime as dt
    if body.block not in ("mot", "expert", "ai"):
        raise HTTPException(422, "block должен быть mot / expert / ai")
    if not (body.category or body.size_key or body.item_name):
        raise HTTPException(422, "укажите категорию, типоразмер или позицию")
    # заменяем прежнюю оценку той же привязки в том же блоке
    q = session.query(ExpertYield).filter(ExpertYield.block == body.block)
    for f, v in (("category", body.category), ("size_key", body.size_key),
                 ("item_name", body.item_name)):
        q = q.filter(getattr(ExpertYield, f) == v)
    q.delete(synchronize_session=False)
    session.add(ExpertYield(
        block=body.block, category=body.category, size_key=body.size_key,
        item_name=body.item_name, good_percent=body.good_percent,
        scrap_percent=100 - body.good_percent, expert_name=body.expert_name,
        confidence="high", approved_by=body.expert_name or user.login,
        approved_at=dt.datetime.now(dt.timezone.utc),
        notes=body.notes or "утверждено через UI"))
    session.commit()
    return {"ok": True}


@app.delete("/api/expert-yields/{yid}")
def delete_expert_yield(yid: int, session: Session = Depends(get_session),
                        user: User = Depends(require_director)):
    from ..db.models import ExpertYield
    session.query(ExpertYield).filter(ExpertYield.id == yid).delete()
    session.commit()
    return {"ok": True}


@app.get("/api/families")
def list_families(session: Session = Depends(get_session)):
    from ..normalize import FAMILY_PATTERNS
    return [f for f, _ in FAMILY_PATTERNS]


@app.get("/api/norms")
def norms(session: Session = Depends(get_session)):
    """Глобальные нормативы (региональные ставки — в /api/norms/regional)."""
    rows = (session.query(ApprovedValue)
            .filter(ApprovedValue.is_current.is_(True),
                    ApprovedValue.scope.is_(None))
            .order_by(ApprovedValue.key).all())
    return [{"key": r.key, "value": r.value, "unit": r.unit, "notes": r.notes,
             "approved_by": r.approved_by} for r in rows]


@app.get("/api/norms/regional")
def regional_norms(session: Session = Depends(get_session)):
    rows = (session.query(ApprovedValue)
            .filter(ApprovedValue.is_current.is_(True),
                    ApprovedValue.scope.isnot(None))
            .order_by(ApprovedValue.value).all())
    return [{"key": r.key, "scope": r.scope, "value": r.value, "unit": r.unit,
             "notes": r.notes} for r in rows]


class NormIn(BaseModel):
    key: str
    value: float
    approved_by: str = "директор"


@app.post("/api/norms")
def set_norm(body: NormIn, session: Session = Depends(get_session),
             user: User = Depends(require_director)):
    session.query(ApprovedValue).filter(
        ApprovedValue.key == body.key,
        ApprovedValue.is_current.is_(True)).update({"is_current": False})
    session.add(ApprovedValue(key=body.key, value=body.value,
                              approved_by=user.login))
    session.commit()
    return {"ok": True}


# ---------------------------------------------------------------- экспорт


@app.get("/api/kp/{doc_id}/export")
def export_kp(doc_id: int, session: Session = Depends(get_session)):
    """Выгрузка в Excel: обе оценки, вес, источники, сопоставление."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    doc = session.get(KpDocument, doc_id)
    if not doc:
        raise HTTPException(404, "КП не найдено")
    wb = Workbook()
    ws = wb.active
    ws.title = "Оценка КП"

    headers = [
        ("Наименование (КП)", 42), ("Сопоставление", 38), ("Кол-во", 10),
        ("Ед.", 7), ("Вес, т", 10),
        ("База: выход %", 12), ("База: источник выхода", 30),
        ("База: реализация ₽", 16), ("База: выкуп ₽", 15),
        ("БП: выход %", 11), ("БП: источник выхода", 30),
        ("БП: реализация ₽", 16), ("БП: выкуп ₽", 15),
        ("Основа лома", 30), ("Комментарий", 24),
    ]
    ws.append([h for h, _ in headers])
    navy = PatternFill("solid", fgColor="1B2A4A")
    for i, (_, width) in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = navy
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32

    money = "#,##0"
    tot = {"w": 0.0, "br": 0.0, "bb": 0.0, "pr": 0.0, "pb": 0.0}
    for p in doc.positions:
        d = _position_payload(session, p, as_of=doc.priced_at)
        w_kg = d["weight_kg"] or d["calc_weight_kg"] or 0
        ws.append([
            d["raw_name"], d["item"]["name"] if d["item"] else "— не сопоставлено",
            d["quantity"], d["unit"], round(w_kg / 1000.0, 3),
            d["base"]["good_percent"], d["base"]["good_source"],
            d["base"]["resale_total"], d["base"]["buyout_total"],
            d["bp"]["good_percent"], d["bp"]["good_source"],
            d["bp"]["resale_total"], d["bp"]["buyout_total"],
            d["bp"]["scrap_basis"] or d["base"]["scrap_basis"], d["comment"],
        ])
        tot["w"] += w_kg / 1000.0
        tot["br"] += d["base"]["resale_total"]
        tot["bb"] += d["base"]["buyout_total"]
        tot["pr"] += d["bp"]["resale_total"]
        tot["pb"] += d["bp"]["buyout_total"]

    last = ws.max_row
    for row in ws.iter_rows(min_row=2, max_row=last):
        for col in (8, 9, 12, 13):
            row[col - 1].number_format = money
        row[4].number_format = "#,##0.000"
        for col in (6, 10):
            row[col - 1].number_format = '0.0"%"'
        for col in (7, 11, 14, 15):
            row[col - 1].alignment = Alignment(wrap_text=True, vertical="top")
            row[col - 1].font = Font(size=9, color="4A5568")

    ws.append(["ИТОГО", "", "", "", round(tot["w"], 3), "", "",
               round(tot["br"], 2), round(tot["bb"], 2), "", "",
               round(tot["pr"], 2), round(tot["pb"], 2), "", ""])
    total_row = ws.max_row
    top = Border(top=Side(style="thin", color="1B2A4A"))
    for i in range(1, len(headers) + 1):
        c = ws.cell(row=total_row, column=i)
        c.font = Font(bold=True)
        c.border = top
        if i in (8, 9, 12, 13):
            c.number_format = money
        if i == 5:
            c.number_format = "#,##0.000"

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last}"

    # второй лист — условия расчёта, чтобы цифры можно было воспроизвести
    info = wb.create_sheet("Условия расчёта")
    info.column_dimensions["A"].width = 34
    info.column_dimensions["B"].width = 52
    rows = [("Документ", doc.title),
            ("Дата оценки (цены на дату)",
             doc.priced_at.strftime("%d.%m.%Y") if doc.priced_at else "—"),
            ("Позиций", len(doc.positions)),
            ("Вес партии, т", round(tot["w"], 3))]
    if doc.file_total_weight_kg:
        file_t = doc.file_total_weight_kg / 1000.0
        diff = tot["w"] - file_t
        rows.append(("Сверка с ИТОГО файла, т",
                     f"{tot['w']:,.3f} против {file_t:,.3f} "
                     f"({'совпадает' if abs(diff) / file_t <= 0.005 else f'расхождение {diff:+,.3f}'})"))
    for norm in (session.query(ApprovedValue)
                 .filter(ApprovedValue.is_current.is_(True),
                         ApprovedValue.scope.is_(None)).all()):
        rows.append((norm.key, f"{norm.value:,.2f} {norm.unit or ''} — {norm.notes or ''}"))
    for sp in (session.query(ScrapPrice)
               .filter(ScrapPrice.is_current.is_(True)).all()):
        rows.append((f"Цена лома: {sp.material}", f"{sp.price_per_tonne:,.0f} ₽/т"))
    for k, v in rows:
        info.append([k, v])
    for row in info.iter_rows():
        row[0].font = Font(bold=True, size=10)
        row[1].alignment = Alignment(wrap_text=True)

    EXPORTS.mkdir(parents=True, exist_ok=True)
    out = EXPORTS / f"export_kp_{doc_id}.xlsx"
    wb.save(out)
    return FileResponse(out, filename=f"Оценка_КП_{doc_id}.xlsx")


@app.get("/api/kp/{doc_id}/offer", response_class=HTMLResponse)
def offer_form(doc_id: int, scenario: str = "base",
               session: Session = Depends(get_session),
               user: User = Depends(current_user)):
    """Печатная форма ценового предложения продавцу (открывается и печатается)."""
    import html as _html

    def money(v: float, digits: int = 0) -> str:
        """Русский формат: неразрывный пробел между разрядами."""
        return f"{v:,.{digits}f}".replace(",", "\u00a0")

    doc = session.get(KpDocument, doc_id)
    if not doc:
        raise HTTPException(404, "КП не найдено")
    if scenario not in ("base", "bp"):
        raise HTTPException(422, "сценарий должен быть base или bp")

    rows, total, weight_t = [], 0.0, 0.0
    for p in doc.positions:
        d = _position_payload(session, p, as_of=doc.priced_at)
        sc = d[scenario]
        w = (d["weight_kg"] or d["calc_weight_kg"] or 0) / 1000.0
        rows.append((d["raw_name"], d["quantity"] or 0, d["unit"] or "", w,
                     sc["buyout_per_unit"] or 0, sc["buyout_total"]))
        total += sc["buyout_total"]
        weight_t += w

    body = "".join(
        f"<tr><td>{i}</td><td>{_html.escape(name)}</td>"
        f"<td class=n>{money(qty)}</td><td>{_html.escape(unit)}</td>"
        f"<td class=n>{money(w, 3)}</td><td class=n>{money(per_unit)}</td>"
        f"<td class=n>{money(value)}</td></tr>"
        for i, (name, qty, unit, w, per_unit, value) in enumerate(rows, 1))
    date = doc.priced_at or doc.created_at
    scen_name = ("базовый (консервативный)" if scenario == "base"
                 else "бизнес-план (оптимистичный)")
    return f"""<!DOCTYPE html><html lang="ru"><head><meta charset="utf-8">
<title>Ценовое предложение — {_html.escape(doc.title or "")}</title><style>
/* документ всегда светлый: его печатают и отправляют продавцу */
:root{{color-scheme:light}}
body{{font-family:-apple-system,Arial,sans-serif;color:#1c2430;background:#fff;
margin:28px;font-size:13px}}
h1{{font-size:19px;color:#1B2A4A;margin:0 0 4px}}
.head{{border-bottom:2px solid #1B2A4A;padding-bottom:10px;margin-bottom:16px}}
.meta{{color:#5a6b82;font-size:12px}}
table{{border-collapse:collapse;width:100%;margin-top:12px}}
th,td{{border:1px solid #c9d3e0;padding:5px 7px;text-align:left}}
th{{background:#EAEFF6;font-size:11px}}
td.n{{text-align:right;font-variant-numeric:tabular-nums;white-space:nowrap}}
tfoot td{{font-weight:700;border-top:2px solid #1B2A4A}}
.sign{{margin-top:34px;display:flex;justify-content:space-between}}
.note{{margin-top:14px;color:#5a6b82;font-size:11px}}
@media print{{.noprint{{display:none}}body{{margin:10px}}}}
</style></head><body>
<div class="head">
  <h1>Ценовое предложение на приобретение имущества</h1>
  <div class="meta">ООО «МетОптТорг», г. Пермь · Перечень: {_html.escape(doc.title or "")}
   · Дата: {date.strftime("%d.%m.%Y") if date else ""} · Вариант расчёта: {scen_name}</div>
</div>
<p>Рассмотрев представленный перечень, предлагаем следующую цену приобретения.
Цены указаны без НДС.</p>
<table><thead><tr><th>№</th><th>Наименование</th><th>Кол-во</th><th>Ед.</th>
<th>Вес, т</th><th>Цена, ₽/ед</th><th>Сумма, ₽</th></tr></thead>
<tbody>{body}</tbody>
<tfoot><tr><td colspan="4">ИТОГО</td><td class="n">{money(weight_t, 3)}</td>
<td></td><td class="n">{money(total)}</td></tr></tfoot></table>
<div class="note">Предложение действительно 14 календарных дней с даты составления.
Окончательная цена уточняется по результатам осмотра и взвешивания.</div>
<div class="sign"><div>Подготовил: {_html.escape(user.full_name or user.login)}</div>
<div>_______________ / _______________</div></div>
<p class="noprint" style="margin-top:24px"><button onclick="window.print()">Печать</button></p>
</body></html>"""


# ---------------------------------------------------------------- UI


@app.get("/", response_class=HTMLResponse)
def index():
    return (WEB / "index.html").read_text(encoding="utf-8")


@app.post("/api/kp/{doc_id}/reprice")
def reprice_kp(doc_id: int, session: Session = Depends(get_session)):
    """Пересчитать КП по текущим ценам (сдвинуть дату оценки на сегодня)."""
    import datetime as dt
    doc = session.get(KpDocument, doc_id)
    if not doc:
        raise HTTPException(404, "КП не найдено")
    doc.priced_at = dt.datetime.now(dt.timezone.utc)
    session.commit()
    return {"ok": True, "priced_at": doc.priced_at.isoformat()}


@app.get("/api/me")
def me(user: User = Depends(current_user)):
    return {"login": user.login, "role": user.role,
            "full_name": user.full_name,
            "can_approve": user.role == "директор"}


# ---------------------------------------------------------------- справочник

OUTLIER_FACTOR = 3.0  # во сколько раз цена должна отличаться от медианы семейства


@app.get("/api/items")
def list_items(q: str | None = None, family: str | None = None,
               flag: str | None = None, page: int = 1, per_page: int = 100,
               session: Session = Depends(get_session)):
    """Актуальные торговые позиции справочника (после дедупликации).

    flag: `merged` — объединённые дубли, `outlier` — ценовые выбросы,
    `priced` — есть цена продаж, `no_price` — цены нет.
    """
    from sqlalchemy import func

    from ..calc.analog_price import family_medians
    from ..db.models import ComponentYield, ItemAlias, PriceQuote
    from ..normalize import normalize_name

    base = session.query(Item).filter(Item.is_trade.is_(True),
                                      Item.is_hidden.is_(False))
    if q:
        base = base.filter(Item.name_normalized.contains(normalize_name(q)))
    if family:
        base = base.filter(Item.family == family)

    # цены и признаки собираем пакетно, не по одной позиции
    alias_counts = dict(session.query(ItemAlias.item_id, func.count(ItemAlias.id))
                        .group_by(ItemAlias.item_id).all())
    if flag == "merged":
        base = base.filter(Item.id.in_(list(alias_counts) or [0]))

    quotes: dict[int, dict[str, PriceQuote]] = {}
    for row in session.query(PriceQuote).filter(
            PriceQuote.quote_type.in_(("sales_fact", "purchase_fact",
                                       "ai_estimate"))).all():
        quotes.setdefault(row.item_id, {})[row.quote_type] = row
    with_composition = {i for (i,) in session.query(ComponentYield.item_id)
                        .filter(ComponentYield.item_id.isnot(None)).distinct()}
    medians = family_medians(session)

    def outlier_of(item, price_row) -> float | None:
        """Во сколько раз цена отличается от медианы семейства (None — норма)."""
        if price_row is None or not item.family:
            return None
        med = medians.get((item.family, (price_row.unit or "").strip().lower()))
        if not med or med <= 0:
            return None
        ratio = price_row.price / med
        return round(ratio, 2) if (ratio >= OUTLIER_FACTOR
                                   or ratio <= 1 / OUTLIER_FACTOR) else None

    rows = base.order_by(Item.name).all()
    enriched = []
    for it in rows:
        q_item = quotes.get(it.id, {})
        sale = q_item.get("sales_fact")
        ratio = outlier_of(it, sale)
        if flag == "outlier" and ratio is None:
            continue
        if flag == "priced" and sale is None:
            continue
        if flag == "no_price" and sale is not None:
            continue
        enriched.append({
            "id": it.id, "name": it.name, "unit": it.unit,
            "family": it.family, "size_key": it.size_key,
            "category": it.category,
            "merged_duplicates": alias_counts.get(it.id, 0),
            "sale_price": sale.price if sale else None,
            "sale_note": sale.notes if sale else None,
            "purchase_price": (q_item["purchase_fact"].price
                               if "purchase_fact" in q_item else None),
            "ai_price": (q_item["ai_estimate"].price
                         if "ai_estimate" in q_item else None),
            "has_composition": it.id in with_composition,
            "outlier_ratio": ratio,
        })
    total = len(enriched)
    start = max(0, (page - 1) * per_page)
    return {"total": total, "page": page, "per_page": per_page,
            "items": enriched[start:start + per_page]}


@app.get("/api/items/{item_id}")
def item_card(item_id: int, session: Session = Depends(get_session)):
    """Карточка позиции: цены, составы, выхода, объединённые дубли."""
    from ..db.models import ComponentYield, ExpertYield, ItemAlias, PriceQuote

    it = session.get(Item, item_id)
    if it is None:
        raise HTTPException(404, "Позиция не найдена")
    prices = (session.query(PriceQuote).filter(PriceQuote.item_id == item_id)
              .order_by(PriceQuote.quoted_at.desc()).all())
    comps = session.query(ComponentYield).filter(
        ComponentYield.item_id == item_id).all()
    # условия строим по непустым полям: сравнение с None даёт IS NULL и
    # притянуло бы чужие категорийные нормативы
    from sqlalchemy import or_

    conds = [ExpertYield.item_id == item_id]
    if it.family:
        conds.append(ExpertYield.category == it.family)
    if it.size_key:
        conds.append(ExpertYield.size_key == it.size_key)
    yields_ = session.query(ExpertYield).filter(or_(*conds)).all()
    aliases = session.query(ItemAlias).filter(ItemAlias.item_id == item_id).all()
    return {
        "id": it.id, "name": it.name, "unit": it.unit, "guid": it.guid,
        "family": it.family, "size_key": it.size_key, "category": it.category,
        "unit_mass_kg": it.unit_mass_kg, "unit_mass_source": it.unit_mass_source,
        "unit_mass_url": it.unit_mass_url,
        "unit_mass_confidence": it.unit_mass_confidence,
        "prices": [{"type": p.quote_type, "price": p.price, "unit": p.unit,
                    "source": p.source, "url": p.source_url,
                    "confidence": p.confidence, "notes": p.notes,
                    "at": p.quoted_at.isoformat() if p.quoted_at else None}
                   for p in prices],
        "compositions": [{"material": c.material, "mass_kg": c.metal_mass_kg,
                          "quantity": c.quantity, "unit": c.unit,
                          "block": c.block, "notes": c.notes} for c in comps],
        "yields": [{"block": y.block, "good_percent": y.good_percent,
                    "scope": ("позиция" if y.item_id else
                              f"типоразмер {y.size_key}" if y.size_key else
                              f"категория {y.category}"),
                    "notes": y.notes} for y in yields_],
        "aliases": [{"guid": a.guid, "name": a.name} for a in aliases],
    }


@app.get("/api/merge-candidates")
def merge_candidates(session: Session = Depends(get_session)):
    """Очередь нечётких дублей: одинаковый набор токенов, разные написания."""
    from ..db.models import MergeCandidate

    rows = (session.query(MergeCandidate)
            .filter(MergeCandidate.status == "pending")
            .order_by(MergeCandidate.id).limit(300).all())
    ids = {r.item_id_a for r in rows} | {r.item_id_b for r in rows}
    items = {i.id: i for i in session.query(Item).filter(Item.id.in_(ids)).all()}
    from ..importers.nomenclature import merge_tier, pick_canonical

    out = []
    for r in rows:
        a, b = items.get(r.item_id_a), items.get(r.item_id_b)
        if a is None or b is None:
            continue
        tier = merge_tier(a, b)
        keep, drop = pick_canonical(session, a, b)
        out.append({"id": r.id, "score": r.score, "tier": tier,
                    "a": {"id": keep.id, "name": keep.name, "unit": keep.unit},
                    "b": {"id": drop.id, "name": drop.name, "unit": drop.unit}})
    return out


@app.post("/api/merge-candidates/{cand_id}/{decision}")
def decide_merge(cand_id: int, decision: str,
                 session: Session = Depends(get_session),
                 user: User = Depends(current_user)):
    """Объединить дубли (guid остаётся в aliases) либо отклонить кандидата."""
    from ..db.models import ItemAlias, MergeCandidate, PriceQuote

    if decision not in ("merge", "reject"):
        raise HTTPException(422, "решение: merge или reject")
    cand = session.get(MergeCandidate, cand_id)
    if cand is None or cand.status != "pending":
        raise HTTPException(404, "Кандидат не найден или уже обработан")
    cand.status = "merged" if decision == "merge" else "rejected"
    cand.decided_by = user.login
    cand.decided_at = dt_now()
    if decision == "merge":
        from ..importers.nomenclature import merge_items, pick_canonical

        a = session.get(Item, cand.item_id_a)
        b = session.get(Item, cand.item_id_b)
        if a is None or b is None:
            raise HTTPException(404, "Позиция не найдена")
        keep, drop = pick_canonical(session, a, b)
        try:
            merge_items(session, keep, drop, user.login)
        except ValueError as e:
            raise HTTPException(422, str(e))
        log_action(session, user, "объединение дублей", "items", keep.id,
                   f"поглощена карточка #{drop.id} «{drop.name}»")
    session.commit()
    return {"ok": True, "decision": decision}


class ItemPatch(BaseModel):
    family: str | None = None
    size_key: str | None = None
    is_trade: bool | None = None
    unit_mass_kg: float | None = None


@app.patch("/api/items/{item_id}")
def patch_item(item_id: int, patch: ItemPatch,
               session: Session = Depends(get_session),
               user: User = Depends(current_user)):
    """Ручная правка карточки: семейство, типоразмер, торговая или нет."""
    it = session.get(Item, item_id)
    if it is None:
        raise HTTPException(404, "Позиция не найдена")
    fields = patch.model_dump(exclude_unset=True)
    for key, value in fields.items():
        setattr(it, key, value)
    if "unit_mass_kg" in fields:
        it.unit_mass_source = "введено вручную"
        it.unit_mass_confidence = "high"
        it.unit_mass_at = dt_now()
        it.unit_mass_url = None
    it.source = "manual"  # ручное решение не перетирается при импорте из 1С
    session.commit()
    log_action(session, user, "правка карточки", "items", it.id,
               "; ".join(f"{k}={v}" for k, v in fields.items()))
    return {"ok": True, "id": it.id, "family": it.family,
            "size_key": it.size_key, "is_trade": it.is_trade}


@app.post("/api/items/reclassify")
def reclassify_items(session: Session = Depends(get_session),
                     user: User = Depends(require_director)):
    """Переклассификация справочника после обновления правил.

    Ручные правки (source='manual') не трогаем — человек уже решил.
    """
    from ..importers.nomenclature import NON_TRADE_NAME_RE
    from ..normalize import classify_family, extract_size_key

    changed_family = hidden = 0
    for it in session.query(Item).filter(Item.source != "manual").all():
        if it.is_trade and NON_TRADE_NAME_RE.search(it.name or ""):
            it.is_trade = False
            hidden += 1
            continue
        if not it.is_trade:
            continue
        family = classify_family(it.name)
        size = extract_size_key(it.name)
        if family != it.family or size != it.size_key:
            it.family, it.size_key = family, size
            changed_family += 1
    session.commit()
    log_action(session, user, "переклассификация справочника", "items", None,
               f"семейств изменено {changed_family}, скрыто непрофильных {hidden}")
    return {"family_updated": changed_family, "moved_to_non_trade": hidden}


class BatchPriceIn(BaseModel):
    document_id: int | None = None   # позиции конкретного КП
    family: str | None = None        # либо всё семейство справочника
    limit: int = 50


@app.post("/api/ai-price/batch")
def ai_price_batch(body: BatchPriceIn, session: Session = Depends(get_session),
                   user: User = Depends(current_user)):
    """Фоновый поиск цен для позиций без цены (КП или семейство справочника)."""
    from .. import ai_price as ai
    from ..db.models import PriceQuote
    from ..db.session import SessionLocal as Factory

    priced = {i for (i,) in session.query(PriceQuote.item_id).filter(
        PriceQuote.quote_type.in_(("sales_fact", "ai_estimate")),
        PriceQuote.item_id.isnot(None)).distinct()}

    q = session.query(Item).filter(Item.is_trade.is_(True),
                                   Item.is_hidden.is_(False))
    if body.document_id is not None:
        doc = session.get(KpDocument, body.document_id)
        if doc is None:
            raise HTTPException(404, "КП не найдено")
        ids = {p.item_id for p in doc.positions if p.item_id}
        q = q.filter(Item.id.in_(ids or [0]))
    elif body.family:
        q = q.filter(Item.family == body.family)
    else:
        raise HTTPException(422, "укажите document_id или family")

    # лом пропускаем сразу: для него цена берётся из прайса, а не из поиска
    candidates = [it.id for it in q.all()
                  if it.id not in priced and it.family != "лом"]
    if not candidates:
        return {"job": None, "message": "все подходящие позиции уже с ценой"}
    try:
        job = ai.start_batch(Factory, candidates[:body.limit], user.login)
    except ai.AiUnavailable as e:
        raise HTTPException(503, str(e))
    log_action(session, user, "пакетный ИИ-поиск цен", "items", None,
               f"позиций в очереди: {job['total']}")
    return {"job": job, "candidates_total": len(candidates)}


@app.get("/api/ai-price/batch/{job_id}")
def ai_price_batch_status(job_id: str):
    from .. import ai_price as ai
    job = ai.get_job(job_id)
    if job is None:
        raise HTTPException(404, "Задача не найдена (сервис мог быть перезапущен)")
    return job


@app.post("/api/items/{item_id}/ai-yield")
def ai_yield_item(item_id: int, session: Session = Depends(get_session),
                  user: User = Depends(current_user)):
    """ИИ-оценка делового выхода для карточки (влияет только на сценарий БП)."""
    from .. import ai_yield as ay
    try:
        result = ay.ai_yield_for_item(session, item_id)
    except ay.AiUnavailable as e:
        raise HTTPException(503, str(e))
    except ValueError as e:
        raise HTTPException(422, str(e))
    if result.get("saved"):
        log_action(session, user, "ИИ-оценка выхода", "items", item_id,
                   f"{result['good_percent']}%")
    return result


@app.post("/api/kp/position/{pos_id}/ai-yield")
def ai_yield_position(pos_id: int, session: Session = Depends(get_session),
                      user: User = Depends(current_user)):
    from .. import ai_yield as ay
    p = session.get(KpPosition, pos_id)
    if not p:
        raise HTTPException(404, "Позиция не найдена")
    if not p.item_id:
        raise HTTPException(422, "Позиция не сопоставлена со справочником")
    try:
        result = ay.ai_yield_for_item(session, p.item_id)
    except ay.AiUnavailable as e:
        raise HTTPException(503, str(e))
    except ValueError as e:
        raise HTTPException(422, str(e))
    return {"result": result, "position": _position_payload(session, p)}


@app.get("/api/logistics-regions")
def logistics_regions(session: Session = Depends(get_session)):
    """Места погрузки с фактической ставкой перевозки (из «Отвесной»)."""
    rows = (session.query(ApprovedValue)
            .filter(ApprovedValue.key == "logistics_cost_per_tonne",
                    ApprovedValue.is_current.is_(True),
                    ApprovedValue.scope.isnot(None))
            .order_by(ApprovedValue.scope).all())
    default = (session.query(ApprovedValue)
               .filter(ApprovedValue.key == "logistics_cost_per_tonne",
                       ApprovedValue.is_current.is_(True),
                       ApprovedValue.scope.is_(None)).first())
    return {"default": {"value": default.value if default else None,
                        "notes": default.notes if default else None},
            "regions": [{"scope": r.scope, "value": r.value, "notes": r.notes}
                        for r in rows]}


class KpPatch(BaseModel):
    seller: str | None = None
    seller_region: str | None = None
    title: str | None = None
    status: str | None = None
    offer_amount: float | None = None
    offer_scenario: str | None = None
    comment: str | None = None


@app.patch("/api/kp/{doc_id}")
def patch_kp(doc_id: int, patch: KpPatch,
             session: Session = Depends(get_session),
             user: User = Depends(current_user)):
    """Правка шапки КП: продавец и место погрузки (влияет на логистику)."""
    doc = session.get(KpDocument, doc_id)
    if doc is None:
        raise HTTPException(404, "КП не найдено")
    fields = patch.model_dump(exclude_unset=True)
    if "status" in fields and fields["status"] not in (
            None, "черновик", "отправлено", "выиграно", "проиграно"):
        raise HTTPException(422, "статус: черновик, отправлено, выиграно или проиграно")
    for key, value in fields.items():
        setattr(doc, key, value)
    # отметки времени проставляем сами — их незачем вводить руками
    if fields.get("status") == "отправлено" and doc.sent_at is None:
        doc.sent_at = dt_now()
    if fields.get("status") in ("выиграно", "проиграно"):
        doc.decided_at = dt_now()
    session.commit()
    log_action(session, user, "правка КП", "kp_documents", doc.id,
               "; ".join(f"{k}={v}" for k, v in fields.items()))
    return {"ok": True, "seller_region": doc.seller_region,
            "seller": doc.seller, "title": doc.title, "status": doc.status,
            "offer_amount": doc.offer_amount}


class AutoMergeIn(BaseModel):
    tier: str = "safe"      # safe — только пунктуация; bu — различие в «б/у»
    apply: bool = False     # без apply возвращается предпросмотр


@app.post("/api/merge-candidates/auto")
def auto_merge(body: AutoMergeIn, session: Session = Depends(get_session),
               user: User = Depends(require_director)):
    """Массовое слияние очевидных дублей. По умолчанию — предпросмотр.

    tier=safe: имена совпадают после удаления пунктуации и пробелов.
    tier=bu:   различие только в признаке «б/у» — требует решения директора.
    """
    from ..db.models import MergeCandidate
    from ..importers.nomenclature import merge_items, merge_tier, pick_canonical

    if body.tier not in ("safe", "bu"):
        raise HTTPException(422, "tier: safe или bu")
    rows = (session.query(MergeCandidate)
            .filter(MergeCandidate.status == "pending").all())
    items = {i.id: i for i in session.query(Item).all()}
    pairs, merged, failed = [], 0, []
    for r in rows:
        a, b = items.get(r.item_id_a), items.get(r.item_id_b)
        if a is None or b is None or merge_tier(a, b) != body.tier:
            continue
        keep, drop = pick_canonical(session, a, b)
        pairs.append({"candidate_id": r.id, "keep": keep.name, "drop": drop.name,
                      "unit": keep.unit})
        if body.apply:
            try:
                merge_items(session, keep, drop, user.login)
            except ValueError as e:
                failed.append({"pair": f"{keep.name} ← {drop.name}",
                               "reason": str(e)})
                continue
            r.status = "merged"
            r.decided_by = user.login
            r.decided_at = dt_now()
            merged += 1
    if body.apply:
        session.commit()
        log_action(session, user, "массовое слияние дублей", "items", None,
                   f"уровень {body.tier}: слито {merged}, отклонено {len(failed)}")
    return {"tier": body.tier, "found": len(pairs), "applied": body.apply,
            "merged": merged, "failed": failed, "preview": pairs[:50]}


@app.post("/api/items/{item_id}/ai-weight")
def ai_weight_item(item_id: int, session: Session = Depends(get_session),
                   user: User = Depends(current_user)):
    """ИИ-поиск массы единицы: в 1С реквизита «Вес» нет."""
    from .. import ai_weight as aw
    try:
        result = aw.ai_weight_for_item(session, item_id)
    except aw.AiUnavailable as e:
        raise HTTPException(503, str(e))
    except ValueError as e:
        raise HTTPException(422, str(e))
    if result.get("saved"):
        log_action(session, user, "ИИ-оценка массы", "items", item_id,
                   f"{result['mass_kg']} кг")
    return result


class WeightBatchIn(BaseModel):
    family: str | None = None
    document_id: int | None = None
    limit: int = 50


@app.post("/api/ai-weight/batch")
def ai_weight_batch(body: WeightBatchIn, session: Session = Depends(get_session),
                    user: User = Depends(current_user)):
    """Фоновый поиск масс для штучных позиций без веса."""
    from .. import ai_price as ai
    from .. import ai_weight as aw
    from ..db.session import SessionLocal as Factory
    from ..normalize import unit_info

    q = session.query(Item).filter(Item.is_trade.is_(True),
                                   Item.is_hidden.is_(False),
                                   Item.unit_mass_kg.is_(None))
    if body.document_id is not None:
        doc = session.get(KpDocument, body.document_id)
        if doc is None:
            raise HTTPException(404, "КП не найдено")
        ids = {p.item_id for p in doc.positions if p.item_id}
        q = q.filter(Item.id.in_(ids or [0]))
    elif body.family:
        q = q.filter(Item.family == body.family)
    else:
        raise HTTPException(422, "укажите document_id или family")

    # масса единицы осмысленна только для штучных позиций
    candidates = [it.id for it in q.all()
                  if (unit_info(it.unit) or ("count",))[0] == "count"]
    if not candidates:
        return {"job": None, "message": "все штучные позиции уже с весом"}
    try:
        job = ai.start_batch(Factory, candidates[:body.limit], user.login,
                             worker=aw.ai_weight_for_item)
    except ai.AiUnavailable as e:
        raise HTTPException(503, str(e))
    log_action(session, user, "пакетный ИИ-поиск масс", "items", None,
               f"позиций в очереди: {job['total']}")
    return {"job": job, "candidates_total": len(candidates)}


@app.get("/api/config")
def client_config():
    """Настройки интерфейса. portal_url — куда ведёт логотип в шапке."""
    import os
    return {"portal_url": os.environ.get("PORTAL_URL") or None,
            "portal_port": os.environ.get("PORTAL_PORT", "8079")}


# ------------------------------------------------------------- рынок лома


@app.post("/api/market-review/upload")
def upload_market_review(file: UploadFile = File(...),
                         session: Session = Depends(get_session),
                         user: User = Depends(current_user)):
    """Еженедельный обзор рынка лома: индексы по регионам и потребителям."""
    from ..importers.market_review import import_market_review

    sf = _store_file(session, file, "market_review")
    try:
        st = import_market_review(session, sf.stored_path, sf.id)
    except (ValueError, KeyError) as e:
        sf.parse_status = "error"
        sf.parse_log = str(e)
        session.commit()
        raise HTTPException(422, f"не удалось разобрать обзор: {e}")
    sf.parse_status = "parsed"
    session.commit()
    log_action(session, user, "импорт обзора рынка", "market_prices", None,
               f"выпуск {st.issue}: регионов {st.regions}, потребителей {st.consumers}")
    return {"issue": st.issue, "regions": st.regions, "consumers": st.consumers,
            "period": [d.isoformat() if d else None for d in (st.period or (None, None))],
            "skipped": st.skipped}


@app.get("/api/market-prices")
def market_prices(kind: str = "region", issue: str | None = None,
                  session: Session = Depends(get_session)):
    from ..calc.market import home_region, latest_issue
    from ..db.models import MarketPrice

    issue = issue or latest_issue(session)
    rows = (session.query(MarketPrice)
            .filter(MarketPrice.issue == issue,
                    MarketPrice.scope_kind == kind)
            .order_by(MarketPrice.price_per_tonne.desc()).all())
    return {"issue": issue, "home_region": home_region(session),
            "prices": [{"scope": r.scope, "price": r.price_per_tonne,
                        "min": r.price_min, "max": r.price_max,
                        "basis": r.basis, "days": r.observations,
                        "period": [r.period_start.isoformat() if r.period_start else None,
                                   r.period_end.isoformat() if r.period_end else None]}
                       for r in rows]}


@app.get("/api/scrap-prices/recommended")
def scrap_recommended(region: str | None = None,
                      session: Session = Depends(get_session)):
    """Рекомендация по каждому материалу: индекс × наш коэффициент."""
    from ..calc.market import recommended_price

    out = {}
    for sp in session.query(ScrapPrice).filter(ScrapPrice.is_current.is_(True)):
        rec = recommended_price(session, sp.material, region)
        if rec:
            out[sp.material] = {**rec, "current": sp.price_per_tonne}
    return out


class RatioIn(BaseModel):
    material: str
    ratio: float
    notes: str = ""


@app.post("/api/scrap-prices/ratio")
def set_market_ratio(body: RatioIn, session: Session = Depends(get_session),
                     user: User = Depends(current_user)):
    from ..calc.market import set_ratio

    if not (0.05 <= body.ratio <= 3):
        raise HTTPException(422, "коэффициент вне разумных пределов 0,05–3")
    set_ratio(session, body.material, body.ratio, body.notes, user.login)
    return {"ok": True}


# ------------------------------------------------------ подбор клиентов


class LeadSearchIn(BaseModel):
    item_id: int | None = None
    query_name: str | None = None
    unit: str | None = None
    quantity: float | None = None


@app.post("/api/leads/search")
def leads_search(body: LeadSearchIn, session: Session = Depends(get_session),
                 user: User = Depends(current_user)):
    """Запустить подбор покупателей под позицию. Отвечает сразу, ищет в фоне."""
    from ..db.models import Item as _Item
    from ..leads import service as leads_service

    name = (body.query_name or "").strip()
    item_id = body.item_id
    if not name and item_id:
        it = session.get(_Item, item_id)
        if it is None:
            raise HTTPException(404, "позиция не найдена")
        name = it.name
    if not name:
        raise HTTPException(422, "укажите позицию или текст запроса")
    if item_id is None:
        # Запрос набран руками — сажаем его на карточку тем же каскадом, что и
        # позиции КП. Без карточки не будет ни единицы учёта, ни семейства,
        # а значит ни цены, ни лидов.
        matched, kind, _ = match_item(session, name, body.unit)
        if matched and kind in ("exact", "alias", "model", "contains"):
            item_id = matched
    camp = leads_service.create(session, item_id=item_id, query_name=name,
                                unit=body.unit, quantity=body.quantity,
                                user=user.login)
    log_action(session, user, "подбор клиентов", "lead_campaigns", camp.id, name)
    leads_service.start(SessionLocal, camp.id)
    return {"campaign_id": camp.id, "status": camp.status}


@app.get("/api/leads/campaigns")
def leads_campaigns(limit: int = 50, session: Session = Depends(get_session)):
    from ..db.models import LeadCampaign
    from ..leads import service as leads_service

    rows = (session.query(LeadCampaign)
            .order_by(LeadCampaign.created_at.desc()).limit(limit).all())
    return [leads_service.campaign_dict(session, c, with_leads=False)
            for c in rows]


@app.get("/api/leads/campaigns/{campaign_id}")
def leads_campaign(campaign_id: int, session: Session = Depends(get_session)):
    from ..db.models import LeadCampaign
    from ..leads import service as leads_service

    camp = session.get(LeadCampaign, campaign_id)
    if camp is None:
        raise HTTPException(404, "кампания не найдена")
    return leads_service.campaign_dict(session, camp)


class LeadUpdateIn(BaseModel):
    status: str | None = None
    comment: str | None = None
    contacts: str | None = None
    owner: str | None = None


LEAD_STATUSES = ("новый", "в работе", "отказ", "сделка")


@app.patch("/api/leads/{lead_id}")
def lead_update(lead_id: int, body: LeadUpdateIn,
                session: Session = Depends(get_session),
                user: User = Depends(current_user)):
    """Отметка менеджера по лиду — сервис не заменяет продавца, а ведёт его."""
    from ..db.models import Lead

    lead = session.get(Lead, lead_id)
    if lead is None:
        raise HTTPException(404, "лид не найден")
    if body.status is not None:
        if body.status not in LEAD_STATUSES:
            raise HTTPException(422, f"статус должен быть одним из: "
                                     f"{', '.join(LEAD_STATUSES)}")
        lead.status = body.status
    if body.comment is not None:
        lead.comment = body.comment
    if body.contacts is not None:
        lead.contacts = body.contacts
    lead.owner = body.owner or user.login
    lead.updated_at = dt_now()
    session.commit()
    log_action(session, user, "правка лида", "leads", lead.id,
               f"статус={lead.status}")
    return {"ok": True, "id": lead.id, "status": lead.status}


@app.get("/api/leads/{lead_id}/draft")
def lead_draft(lead_id: int, session: Session = Depends(get_session)):
    """Черновик письма покупателю. Сервис ничего не отправляет."""
    from ..db.models import Lead
    from ..leads import service as leads_service

    lead = session.get(Lead, lead_id)
    if lead is None:
        raise HTTPException(404, "лид не найден")
    return leads_service.outreach_draft(session, lead.campaign_id, lead.id)


@app.get("/api/leads/campaigns/{campaign_id}/export")
def leads_export(campaign_id: int, session: Session = Depends(get_session)):
    """Выгрузка лидов кампании в Excel — с обоснованием и доказательствами."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    from ..db.models import LeadCampaign
    from ..leads import service as leads_service

    camp = session.get(LeadCampaign, campaign_id)
    if camp is None:
        raise HTTPException(404, "кампания не найдена")
    data = leads_service.campaign_dict(session, camp)

    wb = Workbook()
    ws = wb.active
    ws.title = "Клиенты"
    headers = [("Компания", 40), ("ИНН", 14), ("Канал", 18), ("Родство", 22),
               ("Регион", 18), ("Оценка", 9), ("Ожидаемая цена", 16),
               ("Ед.", 6), ("Обоснование", 70), ("Доказательства", 60),
               ("Наш менеджер", 22), ("Статус", 12), ("Комментарий", 30),
               ("Наш клиент", 12)]
    ws.append([h for h, _ in headers])
    navy = PatternFill("solid", fgColor="1B2A4A")
    for i, (_, width) in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = navy
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for l in data.get("leads", []):
        ev = "; ".join(filter(None, [
            f"{e['title']} {e['url'] or ''}".strip() for e in l["evidence"][:4]]))
        ws.append([l["name"], l["inn"] or "", l["channel"] or "", l["tier"] or "",
                   l["region"] or "", l["score"], l["expected_price"],
                   l["price_unit"] or "", l["why"] or "", ev[:600],
                   l["manager"] or "", l["status"], l["comment"] or "",
                   "да" if l.get("is_existing") else "новый"])
    for row in ws.iter_rows(min_row=2):
        row[8].alignment = Alignment(wrap_text=True, vertical="top")
        row[9].alignment = Alignment(wrap_text=True, vertical="top")

    info = wb.create_sheet("Цена")
    p = data["price"]
    for k, v in (("Позиция", data["query_name"]),
                 ("Единица", p["unit"]),
                 ("Пол по металлу", p["floor"]),
                 ("Ориентир", p["target"]),
                 ("Потолок рынка", p["ceiling"]),
                 ("Основание", p["basis"]),
                 ("Примечание", data.get("error") or "")):
        info.append([k, v])
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 100
    for row in info.iter_rows():
        row[0].font = Font(bold=True, size=10)
        row[1].alignment = Alignment(wrap_text=True)

    EXPORTS.mkdir(parents=True, exist_ok=True)
    out = EXPORTS / f"leads_{campaign_id}.xlsx"
    wb.save(out)
    return FileResponse(out, filename=f"Клиенты_{campaign_id}.xlsx")


@app.get("/api/leads/price-ladder")
def price_ladder_for(item_id: int | None = None, name: str | None = None,
                     unit: str | None = None,
                     session: Session = Depends(get_session)):
    """Разбор цены по ступеням каскада — «откуда взялась цифра»."""
    from ..calc import price_ladder as pl
    from ..db.models import Item as _Item

    item = session.get(_Item, item_id) if item_id else None
    target = (name or (item.name if item else "")).strip()
    if not target:
        raise HTTPException(422, "укажите item_id или name")
    cands = pl.candidates(session, target, unit or (item.unit if item else None),
                          item)
    return {"name": target, "candidates": [c.as_dict() for c in cands]}


# ------------------------------------------------------ мониторинг каналов


class MonitorSourceIn(BaseModel):
    handle: str          # канал, адрес страницы или текст поискового запроса
    kind: str = "telegram"


@app.get("/api/monitor/sources")
def monitor_sources(session: Session = Depends(get_session)):
    from ..db.models import MonitorSource
    from ..monitor import scrapegraph
    from ..monitor import service as mon

    rows = (session.query(MonitorSource)
            .order_by(MonitorSource.is_active.desc(), MonitorSource.handle).all())
    return {
        "interval_min": mon.interval_minutes(),
        "scrapegraph": scrapegraph.is_enabled(),
        "sources": [{
            "id": s.id, "kind": s.kind, "handle": s.handle,
            "target": s.target, "title": s.title,
            "is_active": s.is_active,
            "last_polled_at": (s.last_polled_at.isoformat()
                               if s.last_polled_at else None),
            "last_error": s.last_error,
            "messages_seen": s.messages_seen, "listings_found": s.listings_found,
        } for s in rows]}


@app.post("/api/monitor/sources")
def monitor_add_source(body: MonitorSourceIn,
                       session: Session = Depends(get_session),
                       user: User = Depends(current_user)):
    """Добавить канал. Проверяем читаемость сразу, чтобы не копить пустышки."""
    from ..monitor import service as mon
    from ..monitor import telegram

    from ..monitor import scrapegraph

    if body.kind not in mon.KINDS:
        raise HTTPException(422, f"вид источника должен быть одним из: "
                                 f"{', '.join(mon.KINDS)}")
    try:
        src = mon.add_source(session, body.handle, body.kind, user.login)
    except telegram.TelegramUnavailable as e:
        raise HTTPException(422, str(e))
    except scrapegraph.ScrapeGraphUnavailable as e:
        raise HTTPException(503, str(e))
    except ValueError as e:
        raise HTTPException(422, str(e))
    log_action(session, user, "добавлен источник мониторинга",
               "monitor_sources", src.id, src.handle)
    return {"id": src.id, "handle": src.handle, "kind": src.kind,
            "title": src.title}


class MonitorBulkIn(BaseModel):
    text: str
    kind: str = "telegram"


@app.post("/api/monitor/sources/bulk")
def monitor_add_many(body: MonitorBulkIn, session: Session = Depends(get_session),
                     user: User = Depends(current_user)):
    """Добавить пачку источников и сказать по каждому, читается он или нет."""
    from ..monitor import service as mon

    if body.kind not in mon.KINDS:
        raise HTTPException(422, f"вид источника должен быть одним из: "
                                 f"{', '.join(mon.KINDS)}")
    rows = mon.add_many(session, body.text, body.kind, user.login)
    ok = sum(1 for r in rows if r["ok"])
    log_action(session, user, "массовое добавление источников",
               "monitor_sources", None, f"добавлено {ok} из {len(rows)}")
    return {"added": ok, "total": len(rows), "results": rows}


@app.delete("/api/monitor/sources/{source_id}")
def monitor_drop_source(source_id: int, session: Session = Depends(get_session),
                        user: User = Depends(current_user)):
    from ..db.models import MonitorSource

    src = session.get(MonitorSource, source_id)
    if src is None:
        raise HTTPException(404, "источник не найден")
    src.is_active = False
    session.commit()
    log_action(session, user, "источник мониторинга отключён",
               "monitor_sources", src.id, src.handle)
    return {"ok": True}


@app.post("/api/monitor/poll")
def monitor_poll(source_id: int | None = None,
                 session: Session = Depends(get_session),
                 user: User = Depends(current_user)):
    """Опросить каналы прямо сейчас."""
    from ..monitor import service as mon

    res = mon.poll_all(session, source_id)
    log_action(session, user, "опрос каналов мониторинга", "monitor_sources",
               source_id, f"новых объявлений {res.new_listings}")
    return res.as_dict()


@app.get("/api/monitor/listings")
def monitor_listings(direction: str | None = None, family: str | None = None,
                     keyword: str | None = None,
                     matched: bool = False, limit: int = 100,
                     session: Session = Depends(get_session)):
    from ..db.models import MarketListing
    from ..normalize import web_url as _web_url

    q = session.query(MarketListing).filter(MarketListing.is_hidden.is_(False))
    if direction:
        q = q.filter(MarketListing.direction == direction)
    if family:
        q = q.filter(MarketListing.family == family)
    if keyword:
        q = q.filter(MarketListing.matched_keywords.contains(keyword))
    if matched:
        # «интересное» — распознанная позиция ИЛИ сработавшее ключевое слово:
        # классификатор знает только то, что уже проходило через 1С
        from sqlalchemy import or_
        q = q.filter(or_(MarketListing.family.isnot(None),
                         MarketListing.matched_keywords.isnot(None)))
    rows = (q.order_by(MarketListing.posted_at.desc().nullslast(),
                       MarketListing.id.desc())
            .limit(min(limit, 500)).all())
    return [{
        "id": l.id, "posted_at": l.posted_at.isoformat() if l.posted_at else None,
        "url": _web_url(l.url), "author": l.author, "text": l.text,
        "direction": l.direction, "price": l.price, "price_unit": l.price_unit,
        "family": l.family, "size_key": l.size_key, "item_id": l.item_id,
        "match_kind": l.match_kind, "contacts": l.contacts, "region": l.region,
        "keywords": l.matched_keywords,
    } for l in rows]


class KeywordIn(BaseModel):
    phrase: str
    whole_word: bool = True
    note: str = ""


@app.get("/api/monitor/keywords")
def monitor_keywords(session: Session = Depends(get_session)):
    from ..db.models import MonitorKeyword
    from ..monitor import keywords as kw

    kw.seed_defaults(session)   # пустой список бесполезен, даём стартовый набор
    rows = (session.query(MonitorKeyword)
            .order_by(MonitorKeyword.is_active.desc(),
                      MonitorKeyword.hits.desc(), MonitorKeyword.phrase).all())
    return [{"id": k.id, "phrase": k.phrase, "whole_word": bool(k.whole_word),
             "is_active": bool(k.is_active), "note": k.note, "hits": k.hits or 0,
             "last_hit_at": k.last_hit_at.isoformat() if k.last_hit_at else None}
            for k in rows]


@app.post("/api/monitor/keywords")
def monitor_add_keyword(body: KeywordIn, session: Session = Depends(get_session),
                        user: User = Depends(current_user)):
    from ..db.models import MonitorKeyword
    from ..monitor import keywords as kw

    phrase = (body.phrase or "").strip()
    if len(phrase) < 2:
        raise HTTPException(422, "слишком короткое слово: минимум два символа")
    existing = (session.query(MonitorKeyword)
                .filter(MonitorKeyword.phrase == phrase).first())
    if existing is not None:
        existing.is_active = True
        existing.whole_word = body.whole_word
        session.commit()
        kw.reset_cache()
        return {"id": existing.id, "phrase": existing.phrase}
    row = MonitorKeyword(phrase=phrase, whole_word=body.whole_word,
                         note=body.note or None, added_by=user.login)
    session.add(row)
    session.commit()
    kw.reset_cache()
    log_action(session, user, "добавлено ключевое слово", "monitor_keywords",
               row.id, phrase)
    return {"id": row.id, "phrase": row.phrase}


@app.delete("/api/monitor/keywords/{keyword_id}")
def monitor_drop_keyword(keyword_id: int, session: Session = Depends(get_session),
                         user: User = Depends(current_user)):
    from ..db.models import MonitorKeyword
    from ..monitor import keywords as kw

    row = session.get(MonitorKeyword, keyword_id)
    if row is None:
        raise HTTPException(404, "слово не найдено")
    row.is_active = False
    session.commit()
    kw.reset_cache()
    return {"ok": True}


@app.post("/api/monitor/keywords/rescan")
def monitor_rescan(session: Session = Depends(get_session),
                   user: User = Depends(current_user)):
    """Пересопоставить уже собранные объявления с текущим списком слов.

    Иначе новое слово работало бы только на будущих объявлениях, а всё, что
    канал написал раньше, осталось бы невидимым.
    """
    from ..db.models import MarketListing
    from ..monitor import keywords as kw

    kw.reset_cache()
    touched = 0
    for l in session.query(MarketListing).all():
        hits = kw.match(session, l.text or "")
        value = "; ".join(k.phrase for k in hits) or None
        if value != l.matched_keywords:
            l.matched_keywords = value
            touched += 1
    session.commit()
    log_action(session, user, "пересканирование ленты по словам",
               "market_listings", None, f"изменено {touched}")
    return {"updated": touched}
