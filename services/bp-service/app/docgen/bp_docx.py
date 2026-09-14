"""Бизнес-план в формате Word: сводка по стандарту + P&L."""
from __future__ import annotations

import sqlite3
from pathlib import Path

from docx import Document
from docx.shared import Pt

from .. import calc


def _fmt(v, suffix: str = "") -> str:
    if v is None or v == "":
        return "—"
    if isinstance(v, float):
        return f"{v:,.2f}{suffix}".replace(",", " ")
    return f"{v}{suffix}"


def _kv_table(doc: Document, rows: list[tuple[str, object]]) -> None:
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    for i, (label, value) in enumerate(rows):
        table.rows[i].cells[0].text = label
        table.rows[i].cells[1].text = _fmt(value)
        for p in table.rows[i].cells[0].paragraphs:
            for run in p.runs:
                run.font.bold = True


def build_bp_docx(bp: sqlite3.Row, items: list[sqlite3.Row], costs: list[sqlite3.Row],
                  risks: list[sqlite3.Row], pnl: dict, risk_integral: str,
                  path: Path) -> Path:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10)

    doc.add_heading(f"БИЗНЕС-ПЛАН СДЕЛКИ {bp['bp_number']}", level=0)
    doc.add_paragraph(
        f"МетОптТорг · Стандарт БП v2 · Статус: {bp['status']} · Версия {bp['version']} "
        f"{bp['scenario'] or ''} · Менеджер: {bp['manager'] or '—'}"
    )

    doc.add_heading("1. Стороны сделки и лот", level=1)
    _kv_table(doc, [
        ("Продавец", bp["seller_name"]),
        ("Покупатель", bp["buyer_name"]),
        ("Дивизион", bp["division"]),
        ("№ запроса / перечня", bp["tender_ref"]),
        ("Форма работы", bp["source_type"]),
        ("Стоимость лота БЕЗ НДС (порог), руб", bp["lot_cost"]),
        ("Шаг аукциона, руб", bp["auction_step"]),
        ("Засор чёрного лома, %", bp["contamination_pct"]),
        ("Срок вывоза, мес", bp["removal_months"]),
        ("Месяц начала вывоза", bp["start_month"]),
        ("Вид отгрузки", bp["shipment_type"]),
        ("Перемещение", bp["relocation"]),
        ("Условия оплаты", bp["payment_terms"]),
        ("Особые условия", bp["purchase_special"]),
    ])

    doc.add_heading("2. Состав лота", level=1)
    table = doc.add_table(rows=1, cols=8)
    table.style = "Table Grid"
    for i, h in enumerate(["Поставщик", "Подразделение", "Номенклатура", "Тип закуп.",
                           "Объём, тн", "Объём прод., тн", "ПЛАН прод.",
                           "Цена реализ., руб/тн"]):
        table.rows[0].cells[i].text = h
    for row in pnl["rows"]:
        cells = table.add_row().cells
        for i, v in enumerate([row["supplier"], row["division"], row["nomenclature"],
                               row["purchase_type"], row["volume_t"],
                               row["sale_volume_t"], row["sale_type"],
                               row["sale_price"]]):
            cells[i].text = _fmt(v)
    cells = table.add_row().cells
    cells[0].text = "ИТОГО"
    cells[4].text = _fmt(pnl["purchase_volume"])
    cells[5].text = _fmt(pnl["sale_volume"])

    doc.add_heading("3. Отчёт о прибылях и убытках", level=1)
    _kv_table(doc, [
        ("1. Выручка от реализации без НДС, руб", pnl["revenue"]),
        ("2. Себестоимость закупки (лот) без НДС, руб", pnl["lot_cost"]),
        ("   Закупка с НДС (по типам), руб", pnl["purchase_cost_vat"]),
        ("3. Валовый доход от закупочной стоимости, руб", pnl["gross_purchase"]),
        ("   % наценки к закупке", pnl["markup_pct"]),
        ("4. Переменные расходы, руб", pnl["variable_costs"]),
        ("5. Валовый доход от производственной себестоимости, руб",
         pnl["gross_production"]),
        ("6. Постоянные расходы, руб", pnl["fixed_costs"]),
        ("   в т.ч. персонал (с налогами с ФОТ), руб", pnl["personnel_total"]),
        ("   в т.ч. НДС не возмещённый, руб", pnl["vat_unrecovered"]),
        ("7. Итого прямые затраты, руб", pnl["direct_costs"]),
        ("8. Операционная прибыль, руб", pnl["operating_profit"]),
        ("9. Административные расходы, руб", pnl["admin_costs"]),
        (f"10. Финансовые расходы (капитал {pnl['capital_rate']:.0f}%, "
         f"{pnl['removal_months']:.0f} мес), руб", pnl["capital_cost"]),
        ("11. Прочие доходы и расходы, руб", pnl["other_costs"]),
        ("12. Прибыль до налогообложения, руб", pnl["profit_before_tax"]),
        (f"13. Налог на прибыль ({pnl['tax_rate']:.0f}%), руб", pnl["income_tax"]),
        ("14. ЧИСТАЯ ПРИБЫЛЬ, руб", pnl["net_profit"]),
        ("Чистая прибыль на 1 тн, руб", pnl["net_profit_per_t"]),
        ("Рентабельность продаж, %", pnl["ros_pct"]),
        (f"Порог ({pnl['margin_threshold']:.0f}%)",
         "ВЫШЕ порога" if pnl["above_threshold"] else "НИЖЕ порога"),
    ])

    doc.add_heading("4. Риски", level=1)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for i, h in enumerate(["Тип", "Описание", "Вероятность", "Влияние", "Меры снижения"]):
        table.rows[0].cells[i].text = h
    for rk in risks:
        cells = table.add_row().cells
        # В документ идут слова, а не буквы шкалы: файл читают вне сервиса,
        # и расшифровки «С» и «З» там нет.
        for i, v in enumerate([
                rk["risk_type"], rk["description"],
                calc.risk_label(rk["probability"], calc.PROBABILITY_LABELS),
                calc.risk_label(rk["impact"], calc.IMPACT_LABELS),
                rk["mitigation"]]):
            cells[i].text = _fmt(v)
    _kv_table(doc, [
        ("Интегральная оценка риска", risk_integral),
        ("Экспертная оценка", bp["risk_expert"]),
    ])

    doc.add_heading("5. Итог и рекомендация", level=1)
    _kv_table(doc, [
        ("Рекомендация", bp["recommendation"]),
        ("Комментарий экономиста", bp["economist_comment"]),
        ("Что требует уточнения", bp["clarify_notes"]),
    ])

    doc.add_heading("6. Лист согласования", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for i, role in enumerate(["Менеджер (составитель)", "Экономист", "Руководитель"]):
        cell = table.rows[0].cells[i]
        cell.text = f"{role}\nФИО: ______________\nПодпись: __________\nДата: _____________"

    doc.save(path)
    return path
