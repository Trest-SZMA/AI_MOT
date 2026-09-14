"""Стандартизированный БП в xlsx: листы «Данные» и «P&L» с живыми формулами.

Структура повторяет утверждённый образец BP_1759 (Лукойл-Пермь / Коми):
лист «Данные» — позиции лота с распределением стоимости, параметры сценария
и расчёт привлечённого капитала; лист «P&L» — отчёт о прибылях и убытках,
формулы ссылаются на «Данные».
"""
from __future__ import annotations

import sqlite3
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

F = "Arial"
HDR_FILL = PatternFill("solid", start_color="1F4E5F")
SEC_FILL = PatternFill("solid", start_color="D9E6EC")
CALC_FILL = PatternFill("solid", start_color="FFF7D6")
thin = Side(style="thin", color="B0B0B0")
BORDER = Border(top=thin, bottom=thin, left=thin, right=thin)

HDR = Font(name=F, bold=True, size=9, color="FFFFFF")
TXT = Font(name=F, size=10)
BOLD = Font(name=F, bold=True, size=10)


def _num(v) -> float:
    try:
        return float(v) if v is not None else 0.0
    except (TypeError, ValueError):
        return 0.0


def build_bp_xlsx(bp: sqlite3.Row, items: list[sqlite3.Row],
                  costs: list[sqlite3.Row], pnl: dict, path: Path,
                  schedule: dict | None = None) -> Path:
    wb = Workbook()

    # ─────────────────────────── Лист «Данные» ──────────────────────
    ws = wb.active
    ws.title = "Данные"
    widths = [14, 34, 30, 12, 10, 8, 12, 14, 16, 14, 10, 14, 16, 32, 20, 18, 16]
    for col, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = w

    ws["A1"] = (f"{bp['seller_name'] or 'Продавец не указан'}  |  {bp['bp_number']}"
                f"  |  Сценарий: {bp['scenario'] or '—'}")
    ws["A1"].font = Font(name=F, bold=True, size=12)
    ws["A2"] = (f"Дивизион: {bp['division'] or '—'}  |  {bp['tender_ref'] or ''}  |  "
                f"Покупатель: {bp['buyer_name'] or '—'}")
    ws["A2"].font = Font(name=F, size=9, italic=True)

    headers = ["Поставщик", "Подразделение", "Номенклатура (состав лота)",
               "Категория лома", "Тип покупки", "Ед.изм", "Объём закупки, тн",
               "Цена закупки, руб/тн", "Стоимость закупки БЕЗ НДС, руб",
               "Объём продажи\n(с учётом засора), тн", "ПЛАН\n(тип)",
               "Цена реализации, руб/тн", "Выручка БЕЗ НДС, руб",
               "Номенклатура 1С", "GUID 1С", "Покупатель", "Прибыль позиции, руб"]
    r0 = 4
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=r0, column=col, value=h)
        c.font = HDR
        c.fill = HDR_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = BORDER

    first = r0 + 1
    last = r0 + len(items)
    total_row = last + 1
    # Строки параметров сценария (после итога).
    p_contam = total_row + 3   # засор
    p_lot = p_contam + 1       # порог (стоимость лота)
    p_vat = p_lot + 1
    p_cap = p_vat + 1
    p_year = p_cap + 1
    p_months = p_year + 1
    p_delay = p_months + 1     # отсрочка оплаты покупателем, мес

    for i, it in enumerate(items):
        r = first + i
        vals = [it["supplier"], " / ".join(filter(None, [it["division"], it["warehouse"]])),
                it["nomenclature"], it["category"], it["purchase_type"] or "лом",
                it["unit"] or "тн."]
        for col, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=col, value=v)
            c.font = TXT
            c.border = BORDER
        ws.cell(row=r, column=7, value=_num(it["volume_t"])).border = BORDER
        ws.cell(row=r, column=8,
                value=f"=IFERROR(G${p_lot}/G${total_row},0)").border = BORDER
        ws.cell(row=r, column=9, value=f"=G{r}*H{r}").border = BORDER
        # Засор по типу ПРОДАЖИ (колонка K): труба, проданная ломом, теряет
        # объём так же, как купленный лом (методика БП 1865).
        ws.cell(row=r, column=10,
                value=f'=IF(K{r}="лом",G{r}*(1-G${p_contam}),G{r})').border = BORDER
        ws.cell(row=r, column=11, value=it["sale_type"] or it["purchase_type"]).border = BORDER
        ws.cell(row=r, column=12, value=_num(it["sale_price"])).border = BORDER
        ws.cell(row=r, column=13, value=f"=J{r}*L{r}").border = BORDER
        c14 = ws.cell(row=r, column=14, value=it["nomen_1c"]
                      + (" (подтверждено)" if it["match_confirmed"] else "")
                      if it["nomen_1c"] else "не сопоставлено")
        c14.font = TXT
        c14.border = BORDER
        c15 = ws.cell(row=r, column=15, value=it["nomen_1c_guid"])
        c15.font = Font(name=F, size=8, color="777777")
        c15.border = BORDER
        prow = pnl["rows"][i] if i < len(pnl["rows"]) else {}
        c16 = ws.cell(row=r, column=16, value=it["buyer"])
        c16.font = TXT
        c16.border = BORDER
        c17 = ws.cell(row=r, column=17, value=round(prow.get("profit", 0), 2))
        c17.font = TXT
        c17.border = BORDER
        c17.fill = CALC_FILL
        c17.number_format = "#,##0"
        for col in (8, 9, 10, 13):
            ws.cell(row=r, column=col).fill = CALC_FILL
        for col in (7, 8, 9, 10, 12, 13):
            ws.cell(row=r, column=col).number_format = "#,##0.00"

    ws.cell(row=total_row, column=1, value="ИТОГО").font = BOLD
    for col, formula in [(7, f"=SUM(G{first}:G{last})"),
                         (8, f"=IFERROR(I{total_row}/G{total_row},0)"),
                         (9, f"=SUM(I{first}:I{last})"),
                         (10, f"=SUM(J{first}:J{last})"),
                         (12, f"=IFERROR(M{total_row}/J{total_row},0)"),
                         (13, f"=SUM(M{first}:M{last})")]:
        c = ws.cell(row=total_row, column=col, value=formula if items else 0)
        c.font = BOLD
        c.fill = CALC_FILL
        c.border = BORDER
        c.number_format = "#,##0.00"

    ws.cell(row=total_row + 2, column=1, value="ПАРАМЕТРЫ СЦЕНАРИЯ").font = BOLD
    params = [
        (p_contam, "Засор чёрного лома, %", _num(bp["contamination_pct"]) / 100.0, "0.0%"),
        (p_lot, "Порог закупки (стоимость лота без НДС), руб", _num(bp["lot_cost"]), "#,##0"),
        (p_vat, "Ставка НДС, %", _num(bp["vat_rate"]) / 100.0, "0.0%"),
        (p_cap, "Ставка привлечённого капитала, % годовых",
         _num(bp["capital_rate"]) / 100.0, "0.0%"),
        (p_year, "Кол-во месяцев в году", 12, "0"),
        (p_months, "Кол-во месяцев вывоза/расчёта капитала",
         _num(bp["removal_months"]) or 6, "0.##"),
        (p_delay, "Отсрочка оплаты покупателем, мес",
         _num(pnl.get("payment_delay_months")), "0.##"),
    ]
    for row, label, value, fmt in params:
        ws.cell(row=row, column=1, value=label).font = TXT
        c = ws.cell(row=row, column=7, value=value)
        c.font = BOLD
        c.number_format = fmt

    cap_hdr = p_delay + 2
    ws.cell(row=cap_hdr, column=1,
            value="РАСЧЁТ СТОИМОСТИ ПРИВЛЕЧЁННОГО КАПИТАЛА "
                  "(остаток гасится по мере оплаты)").font = BOLD
    base_row = cap_hdr + 1
    custom_base = _num(bp["capital_base"]) if "capital_base" in bp.keys() else 0.0
    ws.cell(row=base_row, column=1,
            value=("База капитала (задана экономистом), руб" if custom_base
                   else "Стоимость закупки с НДС (расчётная база), руб")).font = TXT
    ws.cell(row=base_row, column=7,
            value=round(custom_base or pnl["purchase_cost_vat"], 2)) \
        .number_format = "#,##0"
    ws.cell(row=base_row, column=8,
            value=("часть лота не кредитуется" if custom_base
                   else "по типам: лом без НДС, труба/прочее +НДС")).font = \
        Font(name=F, size=8, italic=True)
    sch_hdr = base_row + 1
    for col, h in [(1, "Период"), (2, "Привлечённый капитал, руб"),
                   (3, "Длительность, мес"), (7, "Проценты за период, руб")]:
        ws.cell(row=sch_hdr, column=col, value=h).font = BOLD
    # График строится приложением (учитывает отсрочку оплаты и дробный срок);
    # в книге остаются живые формулы процентов от базы и ставки.
    periods = pnl.get("capital_schedule") or []
    base_value = custom_base or pnl["purchase_cost_vat"]
    for i, s in enumerate(periods, start=1):
        r = sch_hdr + i
        ws.cell(row=r, column=1, value=i).font = TXT
        share = (s["balance"] / base_value) if base_value else 0.0
        ws.cell(row=r, column=2, value=f"=G${base_row}*{round(share, 6)}") \
            .number_format = "#,##0"
        ws.cell(row=r, column=3, value=round(s.get("span", 1.0), 4)) \
            .number_format = "0.##"
        ws.cell(row=r, column=7,
                value=f"=B{r}*G${p_cap}/G${p_year}*C{r}").number_format = "#,##0"
    months = max(len(periods), 1)
    cap_total = sch_hdr + months + 1
    ws.cell(row=cap_total, column=1,
            value="ИТОГО стоимость привлечённого капитала, руб").font = BOLD
    c = ws.cell(row=cap_total, column=7,
                value=f"=SUM(G{sch_hdr + 1}:G{sch_hdr + months})")
    c.font = BOLD
    c.fill = CALC_FILL
    c.number_format = "#,##0"

    # ─────────────────────────── Лист «P&L» ─────────────────────────
    ws2 = wb.create_sheet("P&L")
    for col, w in zip("ABCDE", [56, 16, 20, 16, 40]):
        ws2.column_dimensions[col].width = w

    D = "Данные"  # ссылки на лист данных

    def cost_amount(section: str, item: str) -> float:
        return sum(_num(c["amount"]) for c in costs
                   if c["section"] == section and c["item"] == item)

    def section_items(section: str) -> list:
        return [c for c in costs if c["section"] == section]

    r = 1
    ws2.cell(row=r, column=1, value=f"{bp['seller_name'] or ''}  |  {bp['bp_number']}"
             f"  |  Сценарий: {bp['scenario'] or '—'}").font = Font(name=F, bold=True, size=12)
    r += 1
    ws2.cell(row=r, column=1, value=f"Дивизион: {bp['division'] or '—'}  |  "
             f"{bp['tender_ref'] or ''}").font = Font(name=F, size=9, italic=True)
    r += 2

    def line(label, b=None, c=None, d=None, e=None, bold=False, fill=False):
        nonlocal r
        cells = {1: label, 2: b, 3: c, 4: d, 5: e}
        for col, v in cells.items():
            cell = ws2.cell(row=r, column=col, value=v)
            cell.font = BOLD if bold else TXT
            cell.border = BORDER
            if fill and col in (2, 3, 4):
                cell.fill = CALC_FILL
            if col in (2, 3, 4):
                cell.number_format = "#,##0.00" if col != 3 else "#,##0"
        r += 1
        return r - 1

    for col, h in enumerate(["ПОКАЗАТЕЛИ", "Тоннаж / объём", "Суммарно, руб",
                             "На 1 тн, руб", "Комментарий"], start=1):
        c = ws2.cell(row=r, column=col, value=h)
        c.font = HDR
        c.fill = HDR_FILL
        c.border = BORDER
    r += 1

    row_rev = line("1. ВЫРУЧКА ОТ РЕАЛИЗАЦИИ БЕЗ НДС",
                   f"='{D}'!J{total_row}", f"='{D}'!M{total_row}",
                   f"=IFERROR(C{r}/B{r},0)", bold=True, fill=True)
    for t in ("лом", "труба", "цветмет", "кабель", "ДХНО"):
        if pnl["by_type"][t]["volume"] or pnl["by_type"][t]["revenue"]:
            line(f"   {t}",
                 f'=SUMIFS(\'{D}\'!J{first}:J{last},\'{D}\'!K{first}:K{last},"{t}")',
                 f'=SUMIFS(\'{D}\'!M{first}:M{last},\'{D}\'!K{first}:K{last},"{t}")',
                 f"=IFERROR(C{r}/B{r},0)", fill=True)
    row_cost = line("2. СЕБЕСТОИМОСТЬ ЗАКУПКИ БЕЗ НДС (лот)",
                    f"='{D}'!G{total_row}", f"='{D}'!I{total_row}",
                    f"=IFERROR(C{r}/B{r},0)", "Стоимость лота (порог)",
                    bold=True, fill=True)
    line("   Себестоимость закупки с НДС (по типам)", None,
         round(pnl["purchase_cost_vat"], 2), None,
         "лом без НДС; труба/ДХНО/прочее +НДС", fill=True)
    row_gross = line("3. ВАЛОВЫЙ ДОХОД ОТ ЗАКУПОЧНОЙ СТОИМОСТИ", f"=B{row_rev}",
                     f"=C{row_rev}-C{row_cost}", f"=IFERROR(C{r}/B{r},0)",
                     "выручка − себестоимость закупки", bold=True, fill=True)
    line("   % наценки к закупке", None, f"=IFERROR(C{row_gross}/C{row_cost},0)",
         None, fill=True)
    r += 1

    line("4. ПЕРЕМЕННЫЕ (ПРОИЗВОДСТВЕННЫЕ) РАСХОДЫ", bold=True)
    var_first = r
    for c_ in section_items("Переменные"):
        line(f"   {c_['item']}", None, _num(c_["amount"]),
             f"=IFERROR(C{r}/B${row_rev},0)", c_["comment"])
    row_var = line("   ИТОГО переменные расходы", None,
                   f"=SUM(C{var_first}:C{r - 1})" if r > var_first else 0,
                   f"=IFERROR(C{r}/B${row_rev},0)", bold=True, fill=True)
    line("5. ВАЛОВЫЙ ДОХОД ОТ ПРОИЗВОДСТВЕННОЙ СЕБЕСТОИМОСТИ", f"=B{row_rev}",
         f"=C{row_rev}-C{row_cost}-C{row_var}", f"=IFERROR(C{r}/B{r},0)",
         bold=True, fill=True)
    r += 1

    line("6. ПОСТОЯННЫЕ (ПРОИЗВОДСТВЕННЫЕ) РАСХОДЫ", bold=True)
    pers_first = r
    for c_ in section_items("Персонал"):
        line(f"      {c_['item']}", None, _num(c_["amount"]),
             f"=IFERROR(C{r}/B${row_rev},0)", c_["comment"])
    salary_amt = cost_amount("Персонал", "Зарплата")
    payroll_rate = _num(bp["payroll_tax_rate"]) / 100.0
    row_payroll = line("      Налоги с ФОТ", None,
                       round(salary_amt * payroll_rate, 2),
                       f"=IFERROR(C{r}/B${row_rev},0)",
                       f"{payroll_rate:.1%} от статьи «Зарплата»", fill=True)
    row_pers = line("   Расходы на персонал, итого", None,
                    f"=SUM(C{pers_first}:C{row_payroll})",
                    f"=IFERROR(C{r}/B${row_rev},0)", bold=True, fill=True)
    fixed_rows = [row_pers]
    for c_ in section_items("Постоянные"):
        fixed_rows.append(line(f"   {c_['item']}", None, _num(c_["amount"]),
                               f"=IFERROR(C{r}/B${row_rev},0)", c_["comment"]))
    row_vat_unrec = line(
        f"   НДС, не возмещённый {pnl.get('vat_unrecovered_pct', 50):.0f}% "
        "(смена типа при продаже)", None,
        round(pnl["vat_unrecovered"], 2),
        f"=IFERROR(C{r}/B${row_rev},0)", fill=True)
    fixed_rows.append(row_vat_unrec)
    row_fixed = line("   ИТОГО постоянные расходы", None,
                     "=" + "+".join(f"C{x}" for x in fixed_rows),
                     f"=IFERROR(C{r}/B${row_rev},0)", bold=True, fill=True)
    row_direct = line("7. ИТОГО ПРЯМЫЕ (ПРОИЗВОДСТВЕННЫЕ) ЗАТРАТЫ", f"=B{row_rev}",
                      f"=C{row_var}+C{row_fixed}", f"=IFERROR(C{r}/B{r},0)",
                      bold=True, fill=True)
    row_op = line("8. ОПЕРАЦИОННАЯ ПРИБЫЛЬ", f"=B{row_rev}",
                  f"=C{row_rev}-C{row_cost}-C{row_direct}", f"=IFERROR(C{r}/B{r},0)",
                  bold=True, fill=True)
    r += 1

    admin_items = section_items("Административные")
    admin_val = (sum(_num(c_["amount"]) for c_ in admin_items))
    row_admin = line("9. АДМИНИСТРАТИВНЫЕ РАСХОДЫ", None, round(admin_val, 2),
                     f"=IFERROR(C{r}/B${row_rev},0)")
    row_capital = line("10. ФИНАНСОВЫЕ РАСХОДЫ (привлечённый капитал)", None,
                       f"='{D}'!G{cap_total}", f"=IFERROR(C{r}/B${row_rev},0)",
                       f"ставка {_num(bp['capital_rate']):.0f}% годовых, "
                       f"{months} мес", fill=True)
    other_val = sum(_num(c_["amount"]) for c_ in section_items("Прочие"))
    row_other = line("11. ПРОЧИЕ ДОХОДЫ И РАСХОДЫ", None, round(other_val, 2), None)
    row_pbt = line("12. ПРИБЫЛЬ ДО НАЛОГООБЛОЖЕНИЯ", f"=B{row_rev}",
                   f"=C{row_op}-C{row_admin}-C{row_capital}-C{row_other}",
                   f"=IFERROR(C{r}/B{r},0)", bold=True, fill=True)
    tax_rate = _num(bp["tax_rate"]) / 100.0
    # База налога: прибыль до налога + невозмещённый НДС, отнесённый на
    # затраты (он не уменьшает базу) — формула C53 БП 1865.
    row_tax = line(f"13. НАЛОГ НА ПРИБЫЛЬ ({tax_rate:.0%} от базы: прибыль + невозм. НДС)",
                   None,
                   f"=MAX((C{row_pbt}+C{row_vat_unrec})*{tax_rate},0)",
                   f"=IFERROR(C{r}/B${row_rev},0)",
                   fill=True)
    row_np = line("14. ЧИСТАЯ ПРИБЫЛЬ", f"=B{row_rev}", f"=C{row_pbt}-C{row_tax}",
                  f"=IFERROR(C{r}/B{r},0)", bold=True, fill=True)
    r += 1
    line("   Рентабельность продаж (ЧП / Выручка)", None,
         f"=IFERROR(C{row_np}/C{row_rev},0)", None, fill=True)
    line("   Рентабельность операционная (Оп.П / Выручка)", None,
         f"=IFERROR(C{row_op}/C{row_rev},0)", None, fill=True)

    # ─────────────────────────── Лист «График вывоза» ───────────────
    if schedule and schedule["bases"]:
        ws3 = wb.create_sheet("График вывоза")
        ws3.column_dimensions["A"].width = 34
        months = schedule["months"]
        ws3["A1"] = (f"График вывоза {bp['bp_number']} · срок {months} мес · "
                     f"потери при отгрузке {schedule['loss_pct']:.0f}%")
        ws3["A1"].font = Font(name=F, bold=True, size=12)
        r = 3
        for b in schedule["bases"]:
            c = ws3.cell(row=r, column=1,
                         value=f"{b['base']} — {b['volume']:.1f} тн "
                               f"(остаток после плана: {b['left']:.1f})")
            c.font = BOLD
            c.fill = SEC_FILL
            r += 1
            hdr_row = r
            ws3.cell(row=r, column=1, value="Показатель, тн").font = BOLD
            for m in range(1, months + 1):
                cc = ws3.cell(row=r, column=1 + m, value=f"Мес {m}")
                cc.font = BOLD
                cc.fill = SEC_FILL
            r += 1
            for label, key in [("Остаток на начало", "balance_start"),
                               ("Отгрузка (план)", "shipment"),
                               ("Реализация (с потерями)", "sale"),
                               ("Перемещение (справочно)", "relocation"),
                               ("Остаток на конец", "balance_end")]:
                ws3.cell(row=r, column=1, value=label).font = TXT
                for i, row_ in enumerate(b["rows"]):
                    cc = ws3.cell(row=r, column=2 + i, value=round(row_[key], 2))
                    cc.font = TXT
                    cc.number_format = "#,##0.0"
                    if key in ("balance_start", "sale", "balance_end"):
                        cc.fill = CALC_FILL
                r += 1
            r += 1
        ws3.cell(row=r, column=1, value="ИТОГО отгрузка, тн").font = BOLD
        for i, t in enumerate(schedule["month_totals"]):
            ws3.cell(row=r, column=2 + i, value=round(t["shipment"], 2)) \
                .number_format = "#,##0.0"

    wb.save(path)
    return path
