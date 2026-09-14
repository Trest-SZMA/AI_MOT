"""Четыре плоские таблицы для загрузки в 1С и Битрикс24.

Точный состав колонок согласовывается с 1С-специалистом и настройками
импорта Битрикс — правится только в этом модуле.
"""
from __future__ import annotations

import sqlite3
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

F = "Arial"


def _sheet(title: str, headers: list[str], rows: list[list]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    for col, h in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=h).font = Font(name=F, bold=True, size=10)
    for r, row in enumerate(rows, start=2):
        for col, v in enumerate(row, start=1):
            ws.cell(row=r, column=col, value=v).font = Font(name=F, size=10)
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = 24
    return wb


def build_upload_tables(bp: sqlite3.Row, items: list[sqlite3.Row], pnl: dict,
                        costs: list[sqlite3.Row], edges: list[sqlite3.Row],
                        processes: list[sqlite3.Row], out_dir: Path) -> list[Path]:
    n = bp["bp_number"]
    seller, buyer = bp["seller_name"] or "", bp["buyer_name"] or ""
    period = " ".join(filter(None, [bp["start_month"] or "",
                                    f"({bp['removal_months']:g} мес)"
                                    if bp["removal_months"] else ""]))
    paths = []

    # 1. Выручка: план реализации по позициям лота (с номенклатурой 1С).
    n1c = {it["id"]: (it["nomen_1c"], it["nomen_1c_guid"]) for it in items}
    rows = [[n, seller, buyer, r["supplier"] or "", r["division"] or "",
             r["nomenclature"],
             n1c.get(r["id"], (None, None))[0] or "не сопоставлено",
             n1c.get(r["id"], (None, None))[1] or "",
             r["category"] or "", r["sale_type"],
             round(r["sale_volume_t"], 3), round(r["sale_price"], 2),
             round(r["revenue"], 2), period]
            for r in pnl["rows"]]
    wb = _sheet("Выручка", ["№ БП", "Продавец", "Покупатель", "Поставщик",
                            "Подразделение", "Номенклатура продавца",
                            "Номенклатура 1С", "GUID 1С", "Категория",
                            "Тип реализации", "Объём продажи, тн",
                            "Цена реализации, руб/тн", "Выручка без НДС, руб",
                            "Период"], rows)
    p = out_dir / f"{n}_выручка.xlsx"
    wb.save(p)
    paths.append(p)

    # 2. Затраты: статьи P&L + автоматические строки.
    rows = [[n, c["section"], c["item"], round(c["amount"] or 0, 2),
             c["base"] or "", c["comment"] or "", period]
            for c in costs]
    rows += [
        [n, "Закупка", "Себестоимость закупки (лот) без НДС",
         round(pnl["lot_cost"], 2), "", "", period],
        [n, "Персонал", "Налоги с ФОТ (авто)", round(pnl["payroll_tax"], 2), "",
         "от статьи «Зарплата»", period],
        [n, "Постоянные", "НДС не возмещённый (авто)",
         round(pnl["vat_unrecovered"], 2), "", "смена типа при продаже", period],
        [n, "Финансовые", "Стоимость привлечённого капитала (авто)",
         round(pnl["capital_cost"], 2), "",
         f"{pnl['capital_rate']:.0f}% годовых, {pnl['removal_months']:.0f} мес", period],
        [n, "Налоги", "Налог на прибыль (авто)", round(pnl["income_tax"], 2), "",
         f"{pnl['tax_rate']:.0f}%", period],
    ]
    wb = _sheet("Затраты", ["№ БП", "Секция", "Статья затрат", "Сумма, руб",
                            "База/узел", "Комментарий", "Период"], rows)
    p = out_dir / f"{n}_затраты.xlsx"
    wb.save(p)
    paths.append(p)

    # 3. План по транспорту: из графа маршрута (пути по потокам).
    rows = [[n, e["from_label"], e["to_label"], e["flow_group"] or "все потоки",
             e["transport"] or "", round(e["volume_t"] or 0, 3),
             e["comment"] or "", period]
            for e in edges]
    wb = _sheet("План по транспорту",
                ["№ БП", "Откуда", "Куда", "Поток (категория)", "Вид транспорта",
                 "Объём, тн", "Комментарий", "Период"], rows)
    p = out_dir / f"{n}_план_транспорт.xlsx"
    wb.save(p)
    paths.append(p)

    # 4. План по переработке: процессы на узлах (смена номенклатуры).
    rows = [[n, pr["node_label"], pr["work_type"], pr["input_nomen"] or "",
             pr["output_nomen"] or "", round(pr["volume_t"] or 0, 3),
             pr["comment"] or "", period]
            for pr in processes]
    wb = _sheet("План по переработке",
                ["№ БП", "Узел (база/цех/производство)", "Вид работ (1С)",
                 "Входная номенклатура", "Выходная номенклатура", "Объём, тн",
                 "Комментарий", "Период"], rows)
    p = out_dir / f"{n}_план_переработка.xlsx"
    wb.save(p)
    paths.append(p)

    return paths
