"""Импорт справочников из выгрузки 1С (Справочники.xlsx) в базу сервиса.

Запуск:  python import_1c.py [путь к Справочники.xlsx]
Повторный запуск обновляет записи по GUID (без дублей).
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

from app.db import connect, init_db
from app.matcher import normalize

DEFAULT_PATH = "/Users/macpavel/Downloads/Справочники (1).xlsx"


def _s(v) -> str | None:
    if v is None:
        return None
    v = str(v).strip()
    return v or None


def import_all(path: str) -> None:
    init_db()
    conn = connect()
    wb = load_workbook(path, read_only=True, data_only=True)

    # Контрагенты: Контрагент | КонтрагентГуид | Тип
    rows = list(wb["Контрагенты"].iter_rows(values_only=True))[1:]
    for name, guid, ctype in rows:
        if not _s(name):
            continue
        conn.execute(
            "INSERT INTO ref_counterparties (name, guid, ctype) VALUES (?, ?, ?) "
            "ON CONFLICT(guid) DO UPDATE SET name = excluded.name, ctype = excluded.ctype",
            (_s(name), _s(guid), _s(ctype)))
    print(f"Контрагенты: {len(rows)}")

    # Склады: ВысшийРодитель | Родитель | РодительГуид | ... | Склад | СкладГуид | ЭтоГруппа
    rows = list(wb["Склады"].iter_rows(values_only=True))[1:]
    n = 0
    for r in rows:
        top, parent, parent_guid, _, _, name, guid, is_group = r[:8]
        if not _s(name):
            continue
        conn.execute(
            "INSERT INTO ref_warehouses (name, guid, parent_name, parent_guid, top_parent, "
            "is_group) VALUES (?, ?, ?, ?, ?, ?) "
            "ON CONFLICT(guid) DO UPDATE SET name = excluded.name, "
            "parent_name = excluded.parent_name, top_parent = excluded.top_parent, "
            "is_group = excluded.is_group",
            (_s(name), _s(guid), _s(parent), _s(parent_guid), _s(top),
             1 if str(is_group).lower() == "true" else 0))
        n += 1
    print(f"Склады: {n}")

    # Подразделения: Код | Наименование | РодительКод | РодительНаименование
    rows = list(wb["Подразделения"].iter_rows(values_only=True))[1:]
    n = 0
    for code, name, pcode, pname in rows:
        if not _s(name):
            continue
        conn.execute(
            "INSERT INTO ref_divisions (code, name, parent_code, parent_name) "
            "VALUES (?, ?, ?, ?) "
            "ON CONFLICT(code) DO UPDATE SET name = excluded.name, "
            "parent_code = excluded.parent_code, parent_name = excluded.parent_name",
            (_s(code), _s(name), _s(pcode), _s(pname)))
        n += 1
    print(f"Подразделения: {n}")

    # Группы аналитического учёта номенклатуры
    rows = list(wb["ГруппыАналитическогоУчетаНоменк"].iter_rows(values_only=True))[1:]
    n = 0
    for name, guid, parent, parent_guid in rows:
        if not _s(name) or _s(parent) == "НЕ использовать" or _s(name) == "НЕ использовать":
            continue
        conn.execute(
            "INSERT INTO ref_nomen_groups (name, guid, parent_name, parent_guid) "
            "VALUES (?, ?, ?, ?) "
            "ON CONFLICT(guid) DO UPDATE SET name = excluded.name, "
            "parent_name = excluded.parent_name",
            (_s(name), _s(guid), _s(parent), _s(parent_guid)))
        n += 1
    print(f"Группы номенклатуры: {n} (группа «НЕ использовать» пропущена)")

    # Номенклатура 1С (позиции) — лист есть в свежих выгрузках
    if "Номенклатура" in wb.sheetnames:
        it = wb["Номенклатура"].iter_rows(values_only=True)
        header = [str(h) if h else "" for h in next(it)]
        idx = {h: i for i, h in enumerate(header)}
        n = 0
        for r in it:
            name = _s(r[idx["Номенклатура"]])
            if not name:
                continue
            conn.execute(
                "INSERT INTO ref_nomenclature_1c (name, guid, unit, gost, cargo_group, "
                "norm) VALUES (?, ?, ?, ?, ?, ?) "
                "ON CONFLICT(guid) DO UPDATE SET name = excluded.name, "
                "unit = excluded.unit, gost = excluded.gost, "
                "cargo_group = excluded.cargo_group, norm = excluded.norm",
                (name, _s(r[idx["НоменклатураГуид"]]),
                 _s(r[idx["ЕдиницаДляОтчетов"]]) or _s(r[idx["ЕдиницаИзмерения"]]),
                 _s(r[idx["ГОСТ"]]), _s(r[idx["НоменклатурнаяГруппаГрузов"]]),
                 normalize(name)))
            n += 1
        conn.execute("CREATE INDEX IF NOT EXISTS idx_ref_n1c_norm "
                     "ON ref_nomenclature_1c(norm)")
        print(f"Номенклатура 1С (позиции): {n}")

    # Виды работ сотрудников (процессы переработки)
    rows = list(wb["ВидыРаботСотрудников"].iter_rows(values_only=True))[1:]
    for guid, name in rows:
        if not _s(name):
            continue
        conn.execute(
            "INSERT INTO ref_work_types (guid, name) VALUES (?, ?) "
            "ON CONFLICT(guid) DO UPDATE SET name = excluded.name",
            (_s(guid), _s(name)))
    print(f"Виды работ: {len(rows)}")

    # Статьи расходов: Ссылка | СсылкаГуид | Родитель | СчетУчета (колонки по заголовку)
    ws = wb["СтатьиРасходов"]
    it = ws.iter_rows(values_only=True)
    header = [str(h) if h else "" for h in next(it)]
    idx = {h: i for i, h in enumerate(header)}
    n = 0
    for r in it:
        name = _s(r[idx["Ссылка"]])
        if not name:
            continue
        conn.execute(
            "INSERT INTO ref_expense_items (name, guid, parent, account) "
            "VALUES (?, ?, ?, ?) "
            "ON CONFLICT(guid) DO UPDATE SET name = excluded.name, "
            "parent = excluded.parent, account = excluded.account",
            (name, _s(r[idx["СсылкаГуид"]]), _s(r[idx["Родитель"]]),
             _s(r[idx["СчетУчета"]])))
        n += 1
    print(f"Статьи расходов: {n}")

    wb.close()
    conn.commit()
    conn.close()
    print("Импорт завершён.")


if __name__ == "__main__":
    p = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_PATH
    if not Path(p).exists():
        sys.exit(f"Файл не найден: {p}")
    import_all(p)
