# -*- coding: utf-8 -*-
"""Экспорт отчёта за период в Excel и Word."""
import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

BRAND_NAVY = RGBColor(0x14, 0x22, 0x48)

HEADERS = ["№", "Дата постановления", "Номер постановления", "Компания", "Автомобиль",
           "Гос. номер", "Статья КоАП", "Нарушение", "Сумма штрафа, руб.", "Сотрудник",
           "Статус оплаты", "Исполнительные производства", "Примечание"]

COL_WIDTHS = [5, 14, 23, 26, 22, 13, 14, 34, 13, 16, 13, 30, 20]


def _dmy(iso: str) -> str:
    if not iso:
        return ""
    p = iso.split("-")
    return f"{p[2]}.{p[1]}.{p[0]}" if len(p) == 3 else iso


def _ip_cell(ip_docs) -> str:
    if not ip_docs:
        return ""
    lines = []
    for d in ip_docs:
        s = d["doc_type"] or "Документ ФССП"
        if d["ip_number"]:
            s += f" № {d['ip_number']}"
        if d["doc_date"]:
            s += f" от {_dmy(d['doc_date'])}"
        lines.append(s)
    return "\n".join(lines)


def _rows(items):
    rows = []
    for i, item in enumerate(items, 1):
        f, docs = item["fine"], item["ip_docs"]
        rows.append([
            i, _dmy(f["date"]), f["number"] or "", f["company"] or "", f["vehicle"] or "",
            f["plate"] or "", f["article"] or "", f["violation"] or "",
            f["amount"] if f["amount"] is not None else "",
            f["employee"] or "", f["status"], _ip_cell(docs), f["note"] or "",
        ])
    return rows


def _summary_lines(stats, prev=None):
    lines = [
        f"Всего штрафов: {stats['cnt']} на сумму {stats['total']:,.0f} руб. (средний штраф {stats['avg']:,.0f} руб.)",
        f"Оплачено: {stats['paid_cnt']} на сумму {stats['paid_sum']:,.0f} руб. ({stats['paid_pct']:.0f} % суммы)",
        f"Не оплачено: {stats['unpaid_cnt']} на сумму {stats['unpaid_sum']:,.0f} руб. ({stats['unpaid_pct']:.0f} % суммы)",
        f"С исполнительными производствами: {stats['ip_cnt']} на сумму {stats['ip_sum']:,.0f} руб.",
    ]
    if prev:
        dc = stats["cnt"] - prev["stats"]["cnt"]
        ds = stats["total"] - prev["stats"]["total"]
        lines.append(f"Изменение к прошлому месяцу ({prev['name']}): {dc:+d} шт., {ds:+,.0f} руб.")
    return [s.replace(",", " ") for s in lines]


BREAKDOWN_TITLES = [
    ("violations", "По видам нарушений"),
    ("vehicles", "По автомобилям"),
    ("articles", "По статьям"),
    ("companies", "По компаниям"),
    ("employees", "По водителям"),
]


def _stats_sheet(wb, stats, breakdowns, prev, border, head_fill):
    ws = wb.create_sheet("Статистика")
    ws.column_dimensions["A"].width = 48
    for col in "BCD":
        ws.column_dimensions[col].width = 14
    r = 1
    ws.cell(row=r, column=1, value="Сводка за месяц").font = Font(bold=True, size=12, color="142248")
    for line in _summary_lines(stats, prev):
        r += 1
        ws.cell(row=r, column=1, value=line).font = Font(size=10)
    r += 2
    for key, title in BREAKDOWN_TITLES:
        rows = breakdowns.get(key) or []
        if not rows:
            continue
        ws.cell(row=r, column=1, value=title).font = Font(bold=True, size=11, color="142248")
        r += 1
        for c, h in enumerate(["Показатель", "Кол-во, шт.", "Сумма, руб.", "Доля, %"], 1):
            cell = ws.cell(row=r, column=c, value=h)
            cell.font = Font(bold=True, size=10, color="FFFFFF")
            cell.fill = head_fill
            cell.border = border
            cell.alignment = Alignment(horizontal="center")
        for item in rows:
            r += 1
            vals = [item["label"], item["cnt"], item["total"], round(item["share"])]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(row=r, column=c, value=v)
                cell.border = border
                cell.font = Font(size=10)
                cell.alignment = Alignment(horizontal="left" if c == 1 else "right", wrap_text=(c == 1))
        r += 2


def make_xlsx(period_title: str, items, stats, breakdowns=None, prev=None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = period_title[:31]

    thin = Side(style="thin", color="8390B1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    # фирменный тёмно-синий МетОптТорг
    head_fill = PatternFill("solid", fgColor="142248")

    for c, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(bold=True, size=10, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        cell.fill = head_fill
        ws.column_dimensions[get_column_letter(c)].width = COL_WIDTHS[c - 1]

    r = 1
    for row in _rows(items):
        r += 1
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = border
            cell.font = Font(size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=(c in (4, 5, 8, 12, 13)),
                                       horizontal="center" if c in (1, 2, 6, 9, 11) else "left")

    # итоги
    r += 1
    ws.cell(row=r, column=8, value="ИТОГО:").font = Font(bold=True)
    total_cell = ws.cell(row=r, column=9, value=stats["total"])
    total_cell.font = Font(bold=True)
    total_cell.border = border

    r += 2
    for line in _summary_lines(stats, prev):
        ws.cell(row=r, column=2, value=line).font = Font(italic=True, size=10)
        r += 1

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:M{len(items) + 1}"

    if breakdowns:
        _stats_sheet(wb, stats, breakdowns, prev, border, head_fill)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_docx(period_title: str, items, stats, breakdowns=None, prev=None) -> bytes:
    doc = Document()
    for section in doc.sections:
        # альбомная ориентация — таблица широкая
        section.page_width, section.page_height = section.page_height, section.page_width
        section.left_margin = section.right_margin = Cm(1.2)
        section.top_margin = section.bottom_margin = Cm(1.2)

    h = doc.add_heading(f"МетОптТорг — Отчет по штрафам — {period_title}", level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in h.runs:
        run.font.color.rgb = BRAND_NAVY

    for line in _summary_lines(stats, prev):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.add_run(line).italic = True
    doc.add_paragraph()

    table = doc.add_table(rows=1, cols=len(HEADERS))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, name in enumerate(HEADERS):
        hdr[i].text = name
        for par in hdr[i].paragraphs:
            for run in par.runs:
                run.font.bold = True
                run.font.size = Pt(8)

    for row in _rows(items):
        cells = table.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = "" if v is None else (f"{v:,.0f}".replace(",", " ") if isinstance(v, float) else str(v))
            for par in cells[i].paragraphs:
                for run in par.runs:
                    run.font.size = Pt(8)

    row = table.add_row().cells
    row[7].text = "ИТОГО:"
    row[8].text = f"{stats['total']:,.0f}".replace(",", " ")
    for i in (7, 8):
        for par in row[i].paragraphs:
            for run in par.runs:
                run.font.bold = True
                run.font.size = Pt(9)

    if breakdowns:
        h2 = doc.add_heading("Статистика: какие штрафы получаем больше всего", level=2)
        for run in h2.runs:
            run.font.color.rgb = BRAND_NAVY
        for key, title in BREAKDOWN_TITLES:
            rows = breakdowns.get(key) or []
            if not rows:
                continue
            doc.add_paragraph().add_run(title).bold = True
            t = doc.add_table(rows=1, cols=4)
            t.style = "Table Grid"
            for i, name in enumerate(["Показатель", "Кол-во, шт.", "Сумма, руб.", "Доля, %"]):
                t.rows[0].cells[i].text = name
                for par in t.rows[0].cells[i].paragraphs:
                    for run in par.runs:
                        run.font.bold = True
                        run.font.size = Pt(8)
            for item in rows:
                cells = t.add_row().cells
                vals = [item["label"], str(item["cnt"]),
                        f"{item['total']:,.0f}".replace(",", " "), f"{item['share']:.0f}"]
                for i, v in enumerate(vals):
                    cells[i].text = v
                    for par in cells[i].paragraphs:
                        for run in par.runs:
                            run.font.size = Pt(8)
            doc.add_paragraph()

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
